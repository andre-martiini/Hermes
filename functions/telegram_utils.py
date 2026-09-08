"""
Hermes Telegram Utils — helper layer extracted from hermes_core_logic.py.

Contém as funções auxiliares do canal Telegram do Hermes (formatação,
integrações de baixo nível com a API do Telegram, transcrição de
áudio/vídeo, construção de prompts do Gemini, cache de documentos/ações,
etc.) que não são, elas mesmas, os dois manipuladores de topo do webhook
(_process_telegram_message, _handle_telegram_callback) nem os entry points
do Firebase Functions (telegramWebhook, on_telegram_inbound,
on_health_log_red_flag) -- esses continuam em hermes_core_logic.py, que
importa deste módulo o que precisa.

Extraído na íntegra do arquivo original (P02, achado operacional da
sub-entrega 6/N): hermes_core_logic.py sozinho excedia o limite de 200000
caracteres da API de escrita do Argos, bloqueando qualquer shipping futuro
desse arquivo. Split puramente mecânico -- nenhuma linha de lógica foi
reescrita, só movida; ver docs/autonomia/execucao.md para o registro
completo desta sub-entrega.

ATENÇÃO para quem for escrever `mock.patch`/`mock.patch.object` em testes
(achado da revisão adversarial da sub-entrega 9/N): `hermes_core_logic.py`
reexporta todo nome definido aqui (`from telegram_utils import (...)`) por
compatibilidade, então `hermes_core_logic.<nome>` continua existindo como
atributo -- mas isso NÃO redireciona uma chamada de nome livre feita de
DENTRO de uma função que mora neste arquivo. Uma função aqui que chama
`outra_funcao_daqui(...)` (nome livre, sem prefixo de módulo) resolve
`outra_funcao_daqui` pelo namespace DESTE módulo (`telegram_utils`), nunca
pelo de `hermes_core_logic`, mesmo que ambos os módulos tenham um atributo
com esse nome. Um `mock.patch.object(hermes_core_logic, "outra_funcao_daqui", ...)`
não intercepta essa chamada -- é preciso mirar
`mock.patch.object(telegram_utils, "outra_funcao_daqui", ...)`. Isso já
quebrou um teste real nesta sub-entrega (`test_tool_schemas.py`, chamada
interna a `send_message_logged` dentro de `_run_gemini_turn`) e o mesmo
padrão se repete em outras chamadas internas deste arquivo (por exemplo
`_send_telegram_message_with_keyboard`, `_send_telegram_message`,
`generate_content_logged`) -- nenhuma delas tem hoje um teste que dependa de
interceptar a chamada interna, mas se um teste futuro precisar, mire este
módulo, não `hermes_core_logic`.
"""
import copy
import json
import html
import os
import re
import base64
import tempfile
import threading
import time
import unicodedata
import wave
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Optional

import requests as _requests
from firebase_admin import firestore, storage
from google.cloud.firestore_v1 import DocumentReference
from gemini_cost_controls import (
    GEMINI_BALANCED_MODEL,
    GEMINI_TRANSCRIPTION_MODEL,
    GEMINI_TTS_MODEL,
    generate_content_logged,
    send_message_logged,
)

_MAX_HISTORY_TURNS = 20

_MAX_FILE_BYTES = 20 * 1024 * 1024  # 20 MB

_MAX_INLINE_MEDIA_BYTES = 700 * 1024  # base64 stays safely below Firestore field limit

_TTS_MODEL_ID = os.environ.get("TELEGRAM_TTS_MODEL", GEMINI_TTS_MODEL)

_TEXT_MODEL_ID = os.environ.get("TELEGRAM_TEXT_MODEL", GEMINI_BALANCED_MODEL)

_DEFAULT_MALE_VOICE = "Charon"

_MAX_TTS_TRANSCRIPT_CHARS = 1500

_DOC_CACHE: dict = {}

_DOC_CACHE_TTL = 60

_ACTION_SNAPSHOT_CACHE: dict = {}

_ACTION_SNAPSHOT_CACHE_TTL = 45

_TELEGRAM_USER_CACHE: dict = {}

_TELEGRAM_USER_CACHE_TTL = 300

_COPILOT_SESSION_HISTORY_LIMIT = 12

def _cached_doc_get(db, collection: str, document: str):
    key = f"{collection}/{document}"
    now = time.monotonic()
    cached = _DOC_CACHE.get(key)
    if cached and (now - cached[0]) < _DOC_CACHE_TTL:
        return cached[1]
    doc = db.collection(collection).document(document).get()
    _DOC_CACHE[key] = (now, doc)
    return doc

def _safe_telegram_session_id(chat_id: str) -> str:
    safe_chat_id = re.sub(r"[^a-zA-Z0-9_-]+", "_", str(chat_id or "").strip())
    return f"telegram_{safe_chat_id or 'unknown'}"

def _resolve_user_id_for_telegram_chat(db, chat_id: str) -> str | None:
    now = time.monotonic()
    cached = _TELEGRAM_USER_CACHE.get(chat_id)
    if cached and (now - cached[0]) < _TELEGRAM_USER_CACHE_TTL:
        return cached[1]

    user_id = None
    try:
        chat_candidates = [str(chat_id)]
        try:
            chat_candidates.append(int(chat_id))
        except Exception:
            pass
        for field in ("telegram_chat_id", "telegramChatId", "chat_id", "telegram_id"):
            for candidate in chat_candidates:
                docs = (
                    db.collection("usuarios")
                    .where(field, "==", candidate)
                    .limit(1)
                    .get()
                )
                if docs:
                    user_id = docs[0].id
                    break
            if user_id:
                break
    except Exception as exc:
        print(f"[SessionBridge] Falha ao resolver usuario do Telegram chat_id={chat_id}: {exc}")

    _TELEGRAM_USER_CACHE[chat_id] = (now, user_id)
    return user_id

def _ensure_copilot_session(db, chat_id: str, session: dict, first_text: str = "", from_user: dict | None = None) -> str:
    session_id = session.get("copilot_session_id") or _safe_telegram_session_id(chat_id)
    session["copilot_session_id"] = session_id

    session_ref = db.collection("sessoes_copiloto").document(session_id)
    user_id = session.get("userId") or _resolve_user_id_for_telegram_chat(db, chat_id)
    if user_id:
        session["userId"] = user_id

    from_user = from_user or {}
    try:
        snap = session_ref.get()
        base = {
            "channel": "telegram",
            "telegramChatId": str(chat_id),
            "lastMessageAt": firestore.SERVER_TIMESTAMP,
        }
        if user_id:
            base["userId"] = user_id
        if from_user:
            base["telegramUser"] = {
                "id": from_user.get("id"),
                "username": from_user.get("username"),
                "first_name": from_user.get("first_name"),
                "last_name": from_user.get("last_name"),
            }
        if not snap.exists:
            title_seed = (first_text or "Conversa via Telegram").strip().replace("\n", " ")
            base.update({
                "title": (title_seed[:40] + "...") if len(title_seed) > 40 else title_seed,
                "createdAt": firestore.SERVER_TIMESTAMP,
                "taskId": session.get("acao_id") or None,
                "systemId": None,
            })
        session_ref.set(base, merge=True)
    except Exception as exc:
        print(f"[SessionBridge] Falha ao garantir sessao do Copiloto {session_id}: {exc}")

    return session_id

def _persist_copilot_message(
    db,
    session_id: str | None,
    role: str,
    content: str,
    *,
    source: str = "telegram",
    tools_used: list | None = None,
    extra: dict | None = None,
):
    if not session_id or not content:
        return
    try:
        payload = {
            "role": role,
            "content": content,
            "source": source,
            "timestamp": firestore.SERVER_TIMESTAMP,
        }
        if tools_used:
            payload["toolsUsed"] = tools_used
        if extra:
            payload.update(extra)
        db.collection("sessoes_copiloto").document(session_id).collection("mensagens").add(payload)
        db.collection("sessoes_copiloto").document(session_id).set({
            "lastMessageAt": firestore.SERVER_TIMESTAMP,
            "channel": "telegram",
        }, merge=True)
    except Exception as exc:
        print(f"[SessionBridge] Falha ao persistir mensagem na sessao {session_id}: {exc}")

def _sanitize_chat_history(history_contents: list, types) -> list:
    """
    Sanitiza o histórico de conversação para respeitar os requisitos estritos da API Gemini:
    1. Apenas papéis válidos ('user' e 'model')
    2. Remove partes vazias
    3. Mescla mensagens consecutivas com o mesmo papel (evita 'user' seguido de 'user' ou 'model' de 'model')
    4. Deve iniciar com 'user' (descarta leading 'model')
    5. Deve finalizar com 'model' (descarta trailing 'user', já que o próximo chat.send_message envia o novo 'user')
    """
    if not history_contents:
        return []

    cleaned = []
    for c in history_contents:
        role = "model" if getattr(c, "role", None) in ("model", "assistant") else "user"
        parts = []
        for p in getattr(c, "parts", []) or []:
            if hasattr(p, "text") and p.text and p.text.strip():
                parts.append(types.Part(text=p.text))
            elif hasattr(p, "inline_data") and getattr(p.inline_data, "data", None):
                parts.append(p)
            elif hasattr(p, "function_call") and getattr(p.function_call, "name", None):
                parts.append(p)
            elif hasattr(p, "function_response") and getattr(p.function_response, "name", None):
                parts.append(p)

        if not parts:
            continue

        if cleaned and cleaned[-1].role == role:
            cleaned[-1].parts.extend(parts)
        else:
            cleaned.append(types.Content(role=role, parts=parts))

    while cleaned and cleaned[0].role != "user":
        cleaned.pop(0)

    while cleaned and cleaned[-1].role != "model":
        cleaned.pop()

    return cleaned

_SCHEMA_FIELDS_TO_DROP = ("default",)

_BATCH_EDIT_FIELD_SCHEMA = {
    "titulo": {"type": "STRING", "description": "Novo título da ação."},
    "descricao": {"type": "STRING", "description": "Nova descrição."},
    "data_limite": {"type": "STRING", "description": "Data de execução (YYYY-MM-DD)."},
    "data_inicio": {"type": "STRING", "description": "Data de início (YYYY-MM-DD)."},
    "prazo_final": {"type": "STRING", "description": "Prazo final (YYYY-MM-DD)."},
    "horario_inicio": {"type": "STRING", "description": "Horário inicial (HH:MM)."},
    "horario_fim": {"type": "STRING", "description": "Horário final (HH:MM)."},
    "status": {
        "type": "STRING",
        "description": "Status da ação.",
        "enum": ["em andamento", "stand-by", "concluído", "excluído"],
    },
    "tags": {"type": "ARRAY", "items": {"type": "STRING"}, "description": "Lista de tags."},
    "area_tematica": {"type": "STRING", "description": "Área temática válida."},
    "tipo_acao": {"type": "STRING", "description": "Tipo da ação (ex.: fast)."},
    "notas": {"type": "STRING", "description": "Notas da ação."},
}

_TOOL_PARAMETER_OVERRIDES = {
    "editar_acoes_em_lote": {
        "type": "OBJECT",
        "properties": {
            "itens": {
                "type": "ARRAY",
                "description": "Uma entrada por ação a alterar.",
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "task_id": {"type": "STRING", "description": "ID da ação a alterar."},
                        "alteracoes": {
                            "type": "OBJECT",
                            "description": "Somente os campos que mudam nesta ação.",
                            "properties": copy.deepcopy(_BATCH_EDIT_FIELD_SCHEMA),
                        },
                    },
                    "required": ["task_id", "alteracoes"],
                },
            },
            "justificativa": {
                "type": "STRING",
                "description": "Motivo registrado no acompanhamento de cada ação.",
            },
        },
    },
}

def _thinking_config_for_model(types, model_id: str):
    """Devolve a configuração de thinking compatível com a família do modelo.

    Enviar o campo da família errada (`thinking_budget` para um modelo 3.x) faz a
    API recusar a requisição inteira com 400 INVALID_ARGUMENT.
    """
    major = 0
    match = re.search(r"gemini-(\d+)", (model_id or "").lower())
    if match:
        major = int(match.group(1))
    try:
        if major >= 3:
            return types.ThinkingConfig(thinking_level="MINIMAL")
        return types.ThinkingConfig(thinking_budget=0)
    except Exception as exc:
        # SDK sem suporte ao campo: melhor omitir a configuração do que enviar
        # algo que o modelo recusa.
        print(f"[Core] Sem thinking_config para {model_id}: {exc}")
        return None

def _sanitize_tool_schema(node):
    """Remove do schema construções que a API Gemini pode recusar."""
    if not isinstance(node, dict):
        return node

    clean = {k: v for k, v in node.items() if k not in _SCHEMA_FIELDS_TO_DROP}

    props = clean.get("properties")
    if isinstance(props, dict):
        clean["properties"] = {k: _sanitize_tool_schema(v) for k, v in props.items()}
    if isinstance(clean.get("items"), dict):
        clean["items"] = _sanitize_tool_schema(clean["items"])
    for key in ("any_of", "anyOf"):
        if isinstance(clean.get(key), list):
            clean[key] = [_sanitize_tool_schema(v) for v in clean[key]]

    if str(clean.get("type", "")).upper().endswith("OBJECT") and not clean.get("properties"):
        # OBJECT sem properties é o resultado de um parâmetro `dict` sem schema.
        # Trocamos por texto JSON — as ferramentas afetadas aceitam os dois formatos.
        clean.pop("properties", None)
        clean.pop("required", None)
        clean["type"] = "STRING"
        descricao = (clean.get("description") or "").strip()
        clean["description"] = (descricao + " Envie um objeto JSON serializado como texto.").strip()

    return clean

def _required_param_names(fn) -> list:
    """Parâmetros realmente obrigatórios — os que não têm valor padrão.

    O SDK marca como obrigatório todo parâmetro cujo padrão é `None`, o que força
    o modelo a inventar valor para campo opcional (horário, prazo, tags).
    """
    import inspect

    try:
        sig = inspect.signature(fn)
    except (TypeError, ValueError):
        return []
    return [
        nome
        for nome, param in sig.parameters.items()
        if param.default is inspect.Parameter.empty
        and param.kind in (param.POSITIONAL_OR_KEYWORD, param.KEYWORD_ONLY)
    ]

def _build_gemini_tools(client, types, tools_list: list) -> list:
    """Converte os callables em FunctionDeclarations saneadas.

    Em caso de falha na conversão, devolve a lista original de callables — o SDK
    volta a gerar os schemas sozinho e nada pior acontece do que já acontecia.
    """
    declarations = []
    for fn in tools_list:
        nome = getattr(fn, "__name__", str(fn))
        try:
            decl = types.FunctionDeclaration.from_callable(client=client, callable=fn)
            data = decl.model_dump(mode="json", exclude_none=True)
            params = _TOOL_PARAMETER_OVERRIDES.get(nome) or data.get("parameters")
            if params:
                params = _sanitize_tool_schema(copy.deepcopy(params))
                obrigatorios = _required_param_names(fn)
                if obrigatorios:
                    params["required"] = obrigatorios
                else:
                    params.pop("required", None)
                data["parameters"] = params
            declarations.append(types.FunctionDeclaration(**data))
        except Exception as exc:
            print(f"[Core] Falha ao declarar a ferramenta '{nome}': {exc} — usando os schemas automáticos do SDK.")
            return list(tools_list)

    return [types.Tool(function_declarations=declarations)]

def _is_invalid_argument_error(err: Exception) -> bool:
    texto = f"{getattr(err, 'code', '')} {getattr(err, 'status', '')} {err}".upper()
    return "INVALID_ARGUMENT" in texto or "400" in texto

def _load_copilot_session_history(db, session_id: str | None, types, limit: int = _COPILOT_SESSION_HISTORY_LIMIT) -> list:
    if not session_id:
        return []
    try:
        msg_docs = (
            db.collection("sessoes_copiloto").document(session_id)
            .collection("mensagens")
            .order_by("timestamp", direction=firestore.Query.DESCENDING)
            .limit(limit)
            .get()
        )
        history = []
        for doc in reversed(list(msg_docs)):
            data = doc.to_dict() or {}
            content = (data.get("content") or "").strip()
            if not content:
                continue
            role = data.get("role")
            if role == "assistant":
                role = "model"
            if role not in ("user", "model"):
                continue
            history.append(types.Content(role=role, parts=[types.Part(text=content)]))
        return _sanitize_chat_history(history, types)
    except Exception as exc:
        print(f"[SessionBridge] Falha ao carregar historico da sessao {session_id}: {exc}")
        return []

def _get_hermes_storage_bucket():
    project_id = (
        os.environ.get("GCLOUD_PROJECT")
        or os.environ.get("GCP_PROJECT")
        or os.environ.get("GOOGLE_CLOUD_PROJECT")
        or "gestao-hermes"
    )
    candidates = [
        os.environ.get("FIREBASE_STORAGE_BUCKET"),
        f"{project_id}.firebasestorage.app",
        f"{project_id}.appspot.com",
        os.environ.get("TELEGRAM_MEDIA_BUCKET"),
        os.environ.get("HERMES_STORAGE_BUCKET"),
        f"{project_id}-telegram-media-us-central1",
        f"{project_id}-slides-us-central1",
    ]

    checked = []
    for bucket_name in [name for name in candidates if name]:
        bucket = storage.bucket(bucket_name)
        checked.append(bucket_name)
        try:
            if bucket.exists():
                return bucket
        except Exception as exc:
            print(f"[Storage] Falha ao verificar bucket {bucket_name}: {exc}")

    raise RuntimeError(
        "Nenhum bucket de Storage disponivel para midias do Telegram. "
        f"Buckets testados: {', '.join(checked)}"
    )

def _format_web_copilot_text_for_telegram(text: str) -> str:
    safe = html.escape(text or "")
    safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe, flags=re.DOTALL)
    safe = re.sub(r"__(.+?)__", r"<b>\1</b>", safe, flags=re.DOTALL)
    safe = re.sub(r"\*(.+?)\*", r"<i>\1</i>", safe, flags=re.DOTALL)
    safe = re.sub(r"\[(.*?)\]\((task:[^)]+)\)", r"\1 (\2)", safe)
    safe = re.sub(r"\[(.*?)\]\((tool:diagnosis:[^)]+)\)", r"\1 (\2)", safe)
    return safe.strip()

def _call_web_callable(
    *,
    function_name: str,
    data: dict,
    user_uid: str | None = None,
    timeout: int = 480,
) -> dict:
    project_id = (
        os.environ.get("GCLOUD_PROJECT")
        or os.environ.get("GCP_PROJECT")
        or os.environ.get("GOOGLE_CLOUD_PROJECT")
        or "gestao-hermes"
    )
    region = os.environ.get("FUNCTION_REGION") or os.environ.get("GOOGLE_CLOUD_REGION") or "us-central1"
    env_key = re.sub(r"[^A-Z0-9]+", "_", function_name.upper()) + "_URL"
    legacy_env_key = "ASK_COPILOTO_HERMES_URL" if function_name == "askCopilotoHermes" else None
    url = (
        (os.environ.get(legacy_env_key) if legacy_env_key else None)
        or os.environ.get(env_key)
        or f"https://{region}-{project_id}.cloudfunctions.net/{function_name}"
    )

    payload = {"data": data}
    headers = {"Content-Type": "application/json"}
    if user_uid:
        headers["X-Hermes-User-Uid"] = user_uid

    resp = _requests.post(url, json=payload, headers=headers, timeout=timeout)
    if not resp.ok:
        raise RuntimeError(f"{function_name} HTTP {resp.status_code}: {resp.text[:300]}")

    response_data = resp.json()
    if response_data.get("error"):
        raise RuntimeError(json.dumps(response_data["error"], ensure_ascii=False))
    result = response_data.get("result", response_data)
    return result if isinstance(result, dict) else {"result": str(result or "")}

def _run_web_copilot_engine(
    *,
    prompt: str,
    session_id: str,
    task_id: str | None = None,
    system_id: str | None = None,
    user_uid: str | None = None,
    drive_files: list[dict] | None = None,
) -> dict:
    return _call_web_callable(
        function_name="askCopilotoHermes",
        data={
            "prompt": prompt,
            "sessionId": session_id,
            "taskId": task_id,
            "systemId": system_id,
            "routingIndex": [],
            "driveFiles": drive_files or [],
        },
        user_uid=user_uid,
        timeout=480,
    )

def _upload_telegram_file_to_drive(file_bytes: bytes, file_name: str, mime_type: str) -> dict | None:
    """Sobe um arquivo recebido no Telegram para o Google Drive, no mesmo formato
    que o Copiloto Hermes (askCopilotoHermes) espera em driveFiles — permitindo que
    arquivos enviados via Telegram sejam indexados no RAG e vinculados à ação
    (pool_dados) exatamente como no upload feito pela interface web."""
    try:
        import io as _io
        from googleapiclient.http import MediaIoBaseUpload
        from main import get_drive_service
        service = get_drive_service()
        file_metadata = {'name': file_name}
        media = MediaIoBaseUpload(_io.BytesIO(file_bytes), mimetype=mime_type, resumable=True)
        uploaded = service.files().create(body=file_metadata, media_body=media, fields='id, webViewLink').execute()
        drive_file_id = uploaded.get('id')
        if not drive_file_id:
            return None
        try:
            service.permissions().create(fileId=drive_file_id, body={'type': 'anyone', 'role': 'reader'}).execute()
        except Exception as perm_e:
            print(f"[Telegram] Aviso: não foi possível tornar o arquivo público no Drive: {perm_e}")
        return {"driveFileId": drive_file_id, "driveFileName": file_name}
    except Exception as exc:
        print(f"[Telegram] Falha ao subir arquivo para o Drive: {exc}")
        return None

def _extract_task_ids_from_text(text: str) -> list[str]:
    seen = set()
    task_ids = []
    for match in re.finditer(r"\btask:([a-zA-Z0-9_-]{8,})", text or ""):
        task_id = match.group(1).strip()
        if task_id not in seen:
            seen.add(task_id)
            task_ids.append(task_id)
    return task_ids

def _find_latest_copilot_card_message_id(db, session_id: str | None, field_name: str) -> str | None:
    if not session_id:
        return None
    try:
        docs = (
            db.collection("sessoes_copiloto").document(session_id)
            .collection("mensagens")
            .order_by("timestamp", direction=firestore.Query.DESCENDING)
            .limit(12)
            .get()
        )
        for doc in docs:
            data = doc.to_dict() or {}
            if data.get(field_name):
                return doc.id
    except Exception as exc:
        print(f"[TelegramCards] Falha ao localizar card {field_name}: {exc}")
    return None

def _set_latest_copilot_card_status(db, session_id: str | None, field_name: str, status: str):
    message_id = _find_latest_copilot_card_message_id(db, session_id, field_name)
    if not message_id:
        return
    try:
        db.collection("sessoes_copiloto").document(session_id).collection("mensagens").document(message_id).update({
            f"{field_name}.status": status,
            f"{field_name}.status_ui": status,
        })
    except Exception as exc:
        print(f"[TelegramCards] Falha ao atualizar status {field_name}: {exc}")

def _summarize_pending_edit(pending_edit: dict) -> str:
    titulo = html.escape(str(pending_edit.get("titulo") or pending_edit.get("task_id") or "acao"))
    lines = ["", "<b>Card de edicao pendente</b>", f"Acao: {titulo}"]
    alteracoes = pending_edit.get("alteracoes") or {}
    for campo, change in list(alteracoes.items())[:5]:
        change = change if isinstance(change, dict) else {}
        original = html.escape(str(change.get("original") or "-"))
        novo = html.escape(str(change.get("novo") or change.get("novo_raw") or "-"))
        lines.append(f"- {html.escape(str(campo))}: {original} -> {novo}")
    if len(alteracoes) > 5:
        lines.append(f"- mais {len(alteracoes) - 5} alteracao(oes)")
    lines.append("No Telegram, use os botoes abaixo para confirmar ou cancelar.")
    return "\n".join(lines)

def _summarize_pending_batch_reschedule(batch_reschedule: dict) -> str:
    items = batch_reschedule.get("items") or []
    lines = ["", "<b>Card de reagendamento em lote</b>", f"Acoes afetadas: {len(items)}"]
    justificativa = batch_reschedule.get("justificativa")
    if justificativa:
        lines.append(f"Motivo: {html.escape(str(justificativa)[:300])}")
    for item in items[:6]:
        titulo = html.escape(str(item.get("titulo") or item.get("task_id") or "acao"))
        original = html.escape(str(item.get("data_limite_original") or "-"))
        nova = html.escape(str(item.get("nova_data_limite") or "-"))
        lines.append(f"- {titulo}: {original} -> {nova}")
    if len(items) > 6:
        lines.append(f"- mais {len(items) - 6} acao(oes)")
    lines.append("No Telegram, use os botoes abaixo para confirmar ou cancelar.")
    return "\n".join(lines)

def _summarize_memory_conflict(conflict: dict) -> str:
    lines = [
        "",
        "<b>Card de conflito de memoria</b>",
        "O Hermes encontrou duas memorias parecidas e precisa de uma decisao.",
    ]
    if conflict.get("existing_text"):
        lines.append(f"Atual: {html.escape(str(conflict.get('existing_text'))[:300])}")
    if conflict.get("proposed_text"):
        lines.append(f"Nova: {html.escape(str(conflict.get('proposed_text'))[:300])}")
    return "\n".join(lines)

def _build_web_copilot_telegram_adaptation(session: dict, delegated: dict, delegated_text: str) -> tuple[str, list | None]:
    rows = []
    extra_lines = []
    pending_cards = session.setdefault("pending_web_cards", {})

    pending_edit = delegated.get("pendingEdit")
    if isinstance(pending_edit, dict) and pending_edit.get("status", "pending") == "pending":
        pending_cards["edit"] = pending_edit
        extra_lines.append(_summarize_pending_edit(pending_edit))
        rows.append([
            {"text": "Confirmar edicao", "callback_data": "webedit_confirm"},
            {"text": "Cancelar", "callback_data": "webedit_cancel"},
        ])

    pending_batch = delegated.get("pendingBatchReschedule")
    if isinstance(pending_batch, dict) and pending_batch.get("status", "pending") == "pending":
        pending_cards["batch_reschedule"] = pending_batch
        extra_lines.append(_summarize_pending_batch_reschedule(pending_batch))
        rows.append([
            {"text": "Confirmar tudo", "callback_data": "webbatch_confirm"},
            {"text": "Cancelar", "callback_data": "webbatch_cancel"},
        ])

    memory_conflict = delegated.get("pendingMemoryConflict")
    if isinstance(memory_conflict, dict):
        pending_cards["memory_conflict"] = memory_conflict
        extra_lines.append(_summarize_memory_conflict(memory_conflict))
        rows.append([
            {"text": "Manter antiga", "callback_data": "webmem_keep_old"},
            {"text": "Manter nova", "callback_data": "webmem_keep_new"},
        ])

    tool_invocation = delegated.get("toolInvocation")
    if isinstance(tool_invocation, dict):
        tool_id = html.escape(str(tool_invocation.get("tool_id") or "ferramenta"))
        extra_lines.append(
            "\n<b>Ferramenta visual</b>\n"
            f"O Copiloto acionou a ferramenta <code>{tool_id}</code>. "
            "No Telegram ainda nao ha uma versao visual completa para esse componente; mantive o resultado textual acima."
        )

    linked_task_ids = _extract_task_ids_from_text(delegated_text)
    for task_id in linked_task_ids[:3]:
        callback_data = f"lock:{task_id}"
        if len(callback_data) <= 64:
            rows.append([{"text": f"Entrar no contexto {task_id[:8]}", "callback_data": callback_data}])

    if pending_cards:
        session["_pending_web_card_updated_at"] = datetime.now(timezone.utc).isoformat()

    return "\n".join(extra_lines), rows or None

def _get_db():
    return firestore.client()

def _perf_now_ms() -> int:
    return int(time.perf_counter() * 1000)

def _perf_mark(perf_state: dict, name: str):
    started_at = perf_state.get("_last_ms", perf_state["start_ms"])
    now_ms = _perf_now_ms()
    perf_state.setdefault("steps", []).append({
        "name": name,
        "duration_ms": max(0, now_ms - started_at),
    })
    perf_state["_last_ms"] = now_ms

def _perf_log(prefix: str, perf_state: dict, extra: dict | None = None):
    payload = {
        "prefix": prefix,
        "total_ms": max(0, _perf_now_ms() - perf_state["start_ms"]),
        "steps": perf_state.get("steps", []),
    }
    if perf_state.get("tool_calls"):
        payload["tool_calls"] = perf_state["tool_calls"]
    if perf_state.get("tool_rounds"):
        payload["tool_rounds"] = perf_state["tool_rounds"]
    if extra:
        payload.update(extra)
    try:
        print(f"[Perf] {json.dumps(payload, ensure_ascii=False)}")
    except Exception:
        print(f"[Perf] {prefix} total_ms={payload['total_ms']}")

def _get_api_keys(db=None):
    db = db or _get_db()
    doc = _cached_doc_get(db, "system", "api_keys")
    return doc.to_dict() or {} if doc.exists else {}

def _get_telegram_token(db=None) -> str:
    keys = _get_api_keys(db)
    token = keys.get("telegram_bot_token") or os.environ.get("TELEGRAM_BOT_TOKEN", "")
    if not token:
        raise RuntimeError("telegram_bot_token não configurado em system/api_keys.")
    return token

def _get_allowed_chat_id() -> Optional[str]:
    return os.environ.get("ALLOWED_TELEGRAM_CHAT_ID")

_HTML_ANCHOR_RE = re.compile(r"<a\s+[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", re.IGNORECASE | re.DOTALL)

def _neutralize_hidden_links(text: str) -> str:
    """Keep URLs visible in Telegram instead of allowing hidden anchor text."""
    def repl(match: re.Match) -> str:
        href = html.unescape((match.group(1) or "").strip())
        label = re.sub(r"<[^>]+>", "", match.group(2) or "").strip() or href
        safe_label = html.escape(label)
        if href.startswith(("http://", "https://", "task:")):
            safe_href = html.escape(href, quote=False)
            return safe_label if href in label else f"{safe_label} ({safe_href})"
        return safe_label

    return _HTML_ANCHOR_RE.sub(repl, text or "")

def _send_telegram_message(token: str, chat_id: str | int, text: str, parse_mode: str = "HTML"):
    """POST direto à Telegram Bot API."""
    text = _neutralize_hidden_links(text)
    # Telegram HTML: truncate at 4096 chars
    if len(text) > 4096:
        text = text[:4090] + "\n..."
    resp = _requests.post(
        f"https://api.telegram.org/bot{token}/sendMessage",
        json={"chat_id": chat_id, "text": text, "parse_mode": parse_mode},
        timeout=30,
    )
    if not resp.ok:
        print(f"[Telegram] sendMessage failed: {resp.status_code} {resp.text[:300]}")
        if parse_mode:
            plain_text = html.unescape(re.sub(r"<[^>]+>", "", text))
            retry_resp = _requests.post(
                f"https://api.telegram.org/bot{token}/sendMessage",
                json={"chat_id": chat_id, "text": plain_text[:4096]},
                timeout=30,
            )
            if not retry_resp.ok:
                print(f"[Telegram] sendMessage plain fallback failed: {retry_resp.status_code} {retry_resp.text[:300]}")
                return None
            try:
                return retry_resp.json().get("result", {}).get("message_id")
            except Exception:
                return None
        return None
    try:
        return resp.json().get("result", {}).get("message_id")
    except Exception:
        return None

def _send_telegram_message_with_keyboard(
    token: str, chat_id: str | int, text: str, inline_keyboard: list, parse_mode: str = "HTML"
):
    """POST sendMessage com InlineKeyboardMarkup."""
    text = _neutralize_hidden_links(text)
    if len(text) > 4096:
        text = text[:4090] + "\n..."
    resp = _requests.post(
        f"https://api.telegram.org/bot{token}/sendMessage",
        json={
            "chat_id": chat_id,
            "text": text,
            "parse_mode": parse_mode,
            "reply_markup": {"inline_keyboard": inline_keyboard},
        },
        timeout=30,
    )
    if not resp.ok:
        print(f"[Telegram] sendMessage+keyboard failed: {resp.status_code} {resp.text[:300]}")
        if parse_mode:
            plain_text = html.unescape(re.sub(r"<[^>]+>", "", text))
            retry_resp = _requests.post(
                f"https://api.telegram.org/bot{token}/sendMessage",
                json={
                    "chat_id": chat_id,
                    "text": plain_text[:4096],
                    "reply_markup": {"inline_keyboard": inline_keyboard},
                },
                timeout=30,
            )
            if not retry_resp.ok:
                print(f"[Telegram] sendMessage+keyboard plain fallback failed: {retry_resp.status_code} {retry_resp.text[:300]}")
                return None
            try:
                return retry_resp.json().get("result", {}).get("message_id")
            except Exception:
                return None
        return None
    try:
        return resp.json().get("result", {}).get("message_id")
    except Exception:
        return None

def _is_action_context_locked(session: dict | None) -> bool:
    session = session or {}
    return session.get("contexto_ativo") == "acao" and bool(session.get("acao_id") or session.get("acao_titulo"))

def _merge_inline_keyboards(inline_keyboard: list | None, trailing_keyboard: list | None) -> list | None:
    keyboards = []
    seen_callbacks = set()

    for keyboard in (inline_keyboard or [], trailing_keyboard or []):
        if not isinstance(keyboard, list):
            continue
        if keyboard and isinstance(keyboard[0], dict):
            keyboard = [keyboard]
        for row in keyboard:
            if not isinstance(row, list):
                continue
            row_copy = []
            for button in row:
                if not isinstance(button, dict):
                    continue
                callback_data = button.get("callback_data")
                if callback_data and callback_data in seen_callbacks:
                    continue
                if callback_data:
                    seen_callbacks.add(callback_data)
                row_copy.append(button)
            if row_copy:
                keyboards.append(row_copy)

    return keyboards or None

def _send_telegram_session_message(
    db,
    token: str,
    chat_id: str | int,
    text: str,
    session: dict | None = None,
    inline_keyboard: list | None = None,
    parse_mode: str = "HTML",
):
    session = session or _get_session(db, str(chat_id))
    final_text = text
    final_keyboard = inline_keyboard

    if _is_action_context_locked(session):
        acao_titulo = session.get("acao_titulo") or session.get("acao_id") or "acao"
        if not final_text.lstrip().startswith("<b>[Contexto:"):
            final_text = f"<b>[Contexto: {acao_titulo}]</b>\n\n{final_text}"
        final_keyboard = _merge_inline_keyboards(final_keyboard, _EXIT_KEYBOARD)

    # --- Medidor de Contexto ---
    ctx = session.get("contexto_ativo", "geral")
    hist_key = "history_acao" if ctx == "acao" else "history"
    raw_history = session.get(hist_key, [])
    # _MAX_HISTORY_TURNS * 2 é o limite de itens (user+model)
    turns = len(raw_history) // 2
    
    if turns >= (_MAX_HISTORY_TURNS * 0.75):  # Começa a avisar a partir de 75% (15 turnos)
        meter_icon = "🔴" if turns >= _MAX_HISTORY_TURNS else "🟡"
        warning = f"\n\n{meter_icon} <b>Contexto: {turns}/{_MAX_HISTORY_TURNS} turnos</b>"
        if turns >= _MAX_HISTORY_TURNS:
            warning += " (Limite atingido, mensagens antigas serão esquecidas. Use /limpar se necessário)."
        # Evita duplicar se já houver o aviso (importante para redifusões ou mensagens de sistema)
        if warning.strip() not in final_text:
            final_text += warning

    if final_keyboard:
        return _send_telegram_message_with_keyboard(token, chat_id, final_text, final_keyboard, parse_mode=parse_mode)
    return _send_telegram_message(token, chat_id, final_text, parse_mode=parse_mode)

def _answer_callback_query(token: str, callback_query_id: str, text: str = None):
    """Confirma o recebimento de um callback_query (remove indicador de loading no Telegram)."""
    payload: dict = {"callback_query_id": callback_query_id}
    if text:
        payload["text"] = text[:200]
    try:
        _requests.post(
            f"https://api.telegram.org/bot{token}/answerCallbackQuery",
            json=payload,
            timeout=5,
        )
    except Exception:
        pass

def _send_telegram_chat_action(token: str, chat_id: str | int, action: str):
    try:
        _requests.post(
            f"https://api.telegram.org/bot{token}/sendChatAction",
            json={"chat_id": chat_id, "action": action},
            timeout=5,
        )
    except Exception:
        pass

def _delete_telegram_message(token: str, chat_id: str | int, message_id: int):
    if not message_id: return
    try:
        _requests.post(
            f"https://api.telegram.org/bot{token}/deleteMessage",
            json={"chat_id": chat_id, "message_id": message_id},
            timeout=5,
        )
    except Exception:
        pass

def _clear_telegram_inline_keyboard(token: str, chat_id: str | int, message_id: int):
    """Remove o teclado inline (botões) de uma mensagem já respondida, para que um duplo
    toque acidental (ou um reenvio de webhook do Telegram) não reexecute o callback."""
    if not message_id:
        return
    try:
        _requests.post(
            f"https://api.telegram.org/bot{token}/editMessageReplyMarkup",
            json={"chat_id": chat_id, "message_id": message_id, "reply_markup": {"inline_keyboard": []}},
            timeout=5,
        )
    except Exception:
        pass

def _send_telegram_typing(token: str, chat_id: str | int):
    _send_telegram_chat_action(token, chat_id, "typing")

@contextmanager
def _telegram_action_heartbeat(token: str, chat_id: str | int, action: str, interval_seconds: int = 4):
    stop_event = threading.Event()

    def _loop():
        while not stop_event.is_set():
            _send_telegram_chat_action(token, chat_id, action)
            stop_event.wait(interval_seconds)

    thread = threading.Thread(target=_loop, daemon=True)
    thread.start()
    try:
        yield
    finally:
        stop_event.set()
        thread.join(timeout=1)

def _send_telegram_voice(token: str, chat_id: str | int, audio_bytes: bytes, filename: str, mime_type: str, caption: str = "") -> bool:
    files = {
        "voice": (filename, audio_bytes, mime_type),
    }
    data = {
        "chat_id": str(chat_id),
    }
    if caption:
        data["caption"] = caption[:1024]
    resp = _requests.post(
        f"https://api.telegram.org/bot{token}/sendVoice",
        data=data,
        files=files,
        timeout=120,
    )
    if not resp.ok:
        print(f"[Telegram] sendVoice failed: {resp.status_code} {resp.text[:500]}")
    return bool(resp.ok)

def _send_telegram_document(token: str, chat_id: str | int, file_bytes: bytes, filename: str, mime_type: str, caption: str = "") -> bool:
    files = {
        "document": (filename, file_bytes, mime_type),
    }
    data = {
        "chat_id": str(chat_id),
    }
    if caption:
        data["caption"] = caption[:1024]
    resp = _requests.post(
        f"https://api.telegram.org/bot{token}/sendDocument",
        data=data,
        files=files,
        timeout=120,
    )
    if not resp.ok:
        print(f"[Telegram] sendDocument failed: {resp.status_code} {resp.text[:500]}")
    return bool(resp.ok)

def _send_telegram_photo(token: str, chat_id: str | int, photo_url: str, caption: str = "") -> bool:
    data = {
        "chat_id": str(chat_id),
        "photo": photo_url,
    }
    if caption:
        data["caption"] = caption[:1024]
    resp = _requests.post(
        f"https://api.telegram.org/bot{token}/sendPhoto",
        json=data,
        timeout=120,
    )
    if not resp.ok:
        print(f"[Telegram] sendPhoto failed: {resp.status_code} {resp.text[:500]}")
    return bool(resp.ok)

def _get_telegram_file(token: str, file_id: str) -> dict:
    """Calls getFile and returns the file metadata dict."""
    resp = _requests.get(
        f"https://api.telegram.org/bot{token}/getFile",
        params={"file_id": file_id},
        timeout=15,
    )
    resp.raise_for_status()
    data = resp.json()
    return data.get("result", {})

def _download_telegram_file(token: str, file_path: str) -> bytes:
    url = f"https://api.telegram.org/file/bot{token}/{file_path}"
    resp = _requests.get(url, timeout=60)
    resp.raise_for_status()
    return resp.content

def _get_session(db, chat_id: str) -> dict:
    doc = db.collection("telegram_sessions").document(chat_id).get()
    return doc.to_dict() or {"chat_id": chat_id, "contexto_ativo": "geral", "history": []}

def _save_session(db, chat_id: str, session: dict):
    session["updated_at"] = firestore.SERVER_TIMESTAMP
    db.collection("telegram_sessions").document(chat_id).set(session)
    copilot_session_id = session.get("copilot_session_id")
    if copilot_session_id:
        try:
            db.collection("sessoes_copiloto").document(copilot_session_id).set({
                "channel": "telegram",
                "telegramChatId": str(chat_id),
                "taskId": session.get("acao_id") or None,
                "lastMessageAt": firestore.SERVER_TIMESTAMP,
            }, merge=True)
        except Exception as exc:
            print(f"[SessionBridge] Falha ao sincronizar metadados da sessao {copilot_session_id}: {exc}")

def _session_matches_processing_state(session: dict, contexto_ativo: str, acao_id: str | None) -> bool:
    current_contexto = session.get("contexto_ativo", "geral")
    if current_contexto != (contexto_ativo or "geral"):
        return False
    if (contexto_ativo or "geral") == "acao":
        return session.get("acao_id") == acao_id
    return True

def _handle_command(text: str, session: dict) -> Optional[str]:
    """
    Returns a reply string if the message is a state-machine command, else None.
    Side-effects: mutates session.
    """
    text = (text or "").strip()
    m = re.match(r"^/contexto\s+(.+)$", text, re.IGNORECASE)
    if m:
        ctx = m.group(1).strip()
        session["contexto_ativo"] = ctx
        session["history"] = []  # fresh history for new context
        return f"Contexto ativo definido como <b>{ctx}</b>. Histórico reiniciado."

    if re.match(r"^/sair$", text, re.IGNORECASE):
        previous = session.get("acao_titulo") or session.get("contexto_ativo", "geral")
        session["contexto_ativo"] = "geral"
        session["history"] = []
        session["acao_id"] = None
        session["acao_titulo"] = None
        session["acao_context_snapshot"] = None
        session["history_acao"] = []
        return f"Saindo do contexto <b>{previous}</b>. Voltando ao modo geral."

    if re.match(r"^/start$", text, re.IGNORECASE):
        return (
            "👋 Olá! Sou o <b>Hermes Copiloto</b>.\n\n"
            "Comandos disponíveis:\n"
            "• <code>/entrar [termo]</code> — busca ações e trava o contexto nelas\n"
            "• <code>/sair</code> — sai do contexto trancado, retorna ao modo geral\n"
            "• <code>/status</code> — mostra o contexto e histórico atuais\n"
            "• <code>caminhada 2.5</code> — registra bloco de esteira em km (aceita min, passos e kcal)\n\n"
            "Envie texto, áudio ou arquivos. Tamanho máximo: 20 MB."
        )

    if re.match(r"^/status$", text, re.IGNORECASE):
        ctx = session.get("contexto_ativo", "geral")
        acao_titulo = session.get("acao_titulo", "")
        hist_key = "history_acao" if ctx == "acao" else "history"
        turns = len(session.get(hist_key, [])) // 2
        ctx_display = f"ação trancada — <b>{acao_titulo}</b>" if ctx == "acao" else f"<b>{ctx}</b>"
        return f"Contexto ativo: {ctx_display}\nTurnos no histórico: {turns}"

    if re.match(r"^/limpar$", text, re.IGNORECASE):
        ctx = session.get("contexto_ativo", "geral")
        hist_key = "history_acao" if ctx == "acao" else "history"
        session[hist_key] = []
        return f"✅ Histórico do contexto <b>{ctx}</b> foi limpo. Podemos recomeçar!"

    return None

_WALK_REGISTER_RE = re.compile(
    r"^(?:hoje\s+)?/?(?:(?:registrar?|registra|registrei|anotar?|anota)\s+)?"
    r"(?:caminhada|caminhei|esteira)\s*[:\-]?\s*(?:de\s+)?"
    r"(?P<dist>\d+(?:[.,]\d+)?)\s*(?:km|quil[oô]metros?)?"
    r"(?P<rest>.*)$",
    re.IGNORECASE,
)

_WALK_MINUTES_RE = re.compile(r"(\d+)\s*min(?:utos?)?\b", re.IGNORECASE)

_WALK_STEPS_RE = re.compile(r"(\d+)\s*passos\b", re.IGNORECASE)

_WALK_KCAL_RE = re.compile(r"(\d+)\s*(?:kcal|calorias?)\b", re.IGNORECASE)

_WALK_FILLER_RE = re.compile(r"\b(?:em|e|com|de|na|no|esteira|hoje|agora)\b", re.IGNORECASE)

def _format_km(value: float) -> str:
    text = f"{value:.2f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")

def _try_register_walk_block(db, text: str) -> Optional[str]:
    """
    Detecta mensagens como "registrar caminhada 2.5", "caminhada 2,5 km 40 min"
    ou "caminhei 1.8" e grava um bloco em health_exercise_logs/{hoje}.
    Retorna o texto de resposta, ou None para deixar a mensagem seguir ao LLM.
    """
    match = _WALK_REGISTER_RE.match((text or "").strip())
    if not match:
        return None

    rest = match.group("rest") or ""
    minutes_match = _WALK_MINUTES_RE.search(rest)
    steps_match = _WALK_STEPS_RE.search(rest)
    kcal_match = _WALK_KCAL_RE.search(rest)

    # Se sobrar conteúdo além dos extras conhecidos, não é um registro direto
    # (ex.: "caminhada 5 vezes por semana faz bem?") — deixa para o LLM.
    leftovers = rest
    for pattern in (_WALK_MINUTES_RE, _WALK_STEPS_RE, _WALK_KCAL_RE):
        leftovers = pattern.sub(" ", leftovers)
    leftovers = _WALK_FILLER_RE.sub(" ", leftovers)
    if re.sub(r"[\s.,;:!\-–—]+", "", leftovers):
        return None

    distance = float(match.group("dist").replace(",", "."))
    if distance <= 0 or distance > 50:
        return (
            "⚠️ Distância fora do esperado. Informe o valor em km, "
            "ex.: <code>caminhada 2.5</code>."
        )

    from zoneinfo import ZoneInfo
    now = datetime.now(ZoneInfo("America/Sao_Paulo"))
    date_key = now.strftime("%Y-%m-%d")

    block = {
        "id": f"walk_{int(now.timestamp() * 1000)}",
        "time": now.strftime("%H:%M"),
        "distance": distance,
        "source": "telegram",
    }
    if minutes_match:
        block["minutes"] = int(minutes_match.group(1))
    if steps_match:
        block["steps"] = int(steps_match.group(1))
    if kcal_match:
        block["calories"] = int(kcal_match.group(1))

    doc_ref = db.collection("health_exercise_logs").document(date_key)
    transaction = db.transaction()

    @firestore.transactional
    def _append_walk_block(transaction, doc_ref, new_block):
        snapshot = doc_ref.get(transaction=transaction)
        blocks = ((snapshot.to_dict() or {}).get("walkBlocks") or []) if snapshot.exists else []
        blocks = [b for b in blocks if isinstance(b, dict)]
        blocks.append(new_block)
        transaction.set(doc_ref, {"walkBlocks": blocks}, merge=True)
        return blocks

    blocks = _append_walk_block(transaction, doc_ref, block)

    total_km = sum(float(b.get("distance") or 0) for b in blocks)

    try:
        settings_doc = db.collection("health_settings").document("config").get()
        settings = settings_doc.to_dict() or {}
    except Exception:
        settings = {}
    # Mesmos padrões e paradigma de nivelamento da UI (DashboardView):
    # abaixo do mínimo ainda não pontua; do mínimo ao ideal o nível progride
    # continuamente (vermelho -> laranja -> amarelo -> verde); do ideal em
    # diante é "lucro".
    minimum_setting = settings.get("walkingMinimumKm")
    ideal_setting = settings.get("walkingIdealKm")
    minimum_km = float(minimum_setting) if minimum_setting is not None else 3.0
    ideal_km = float(ideal_setting) if ideal_setting is not None else 8.0

    if total_km >= ideal_km:
        surplus = total_km - ideal_km
        surplus_text = f" +{_format_km(surplus)} km de lucro!" if surplus > 0 else ""
        progress = f"🟢 Meta ideal de {_format_km(ideal_km)} km atingida!{surplus_text} 🏆"
    elif total_km >= minimum_km:
        t = (total_km - minimum_km) / (ideal_km - minimum_km) if ideal_km > minimum_km else 1.0
        level_emoji = "🔴" if t < 0.25 else "🟠" if t < 0.5 else "🟡" if t < 0.75 else "🟢"
        progress = (
            f"{level_emoji} Mínimo ok — nível {round(t * 100)}% entre o mínimo "
            f"({_format_km(minimum_km)} km) e o ideal ({_format_km(ideal_km)} km). "
            f"Faltam {_format_km(ideal_km - total_km)} km para o ideal."
        )
    else:
        progress = (
            f"⚪ Abaixo do mínimo — faltam {_format_km(minimum_km - total_km)} km "
            f"para entrar na zona de pontuação (mínimo de {_format_km(minimum_km)} km)."
        )

    extras = []
    if block.get("minutes"):
        extras.append(f"{block['minutes']} min")
    if block.get("steps"):
        extras.append(f"{block['steps']} passos")
    if block.get("calories"):
        extras.append(f"{block['calories']} kcal")
    extras_text = f" ({', '.join(extras)})" if extras else ""

    plural = "blocos" if len(blocks) > 1 else "bloco"
    return (
        f"🚶 Caminhada registrada: <b>{_format_km(distance)} km</b> às {block['time']}{extras_text}.\n"
        f"Total de hoje: <b>{_format_km(total_km)} km</b> em {len(blocks)} {plural}.\n"
        f"{progress}"
    )

_DRIVE_ID_PATTERNS = [
    re.compile(r"/d/([a-zA-Z0-9_-]{10,})"),
    re.compile(r"[?&]id=([a-zA-Z0-9_-]{10,})"),
    re.compile(r"/folders/([a-zA-Z0-9_-]{10,})"),
]

def _extract_drive_file_id(value: str | None) -> str | None:
    value = (value or "").strip()
    if re.fullmatch(r"[a-zA-Z0-9_-]{10,}", value):
        return value
    for pattern in _DRIVE_ID_PATTERNS:
        match = pattern.search(value)
        if match:
            return match.group(1)
    return None

def _parse_diary_file_note(note: str | None) -> dict | None:
    """Parse FILE:: diary rich notes created by the web diary UI."""
    note = (note or "").strip()
    if not note.startswith("FILE::"):
        return None

    payload = note[len("FILE::"):]
    name = ""
    value = ""
    if payload.startswith("JSON::"):
        try:
            parsed = json.loads(payload[len("JSON::"):])
            if isinstance(parsed, dict):
                name = str(parsed.get("n") or "")
                value = str(parsed.get("v") or "")
        except Exception:
            return None
    else:
        separator_index = payload.find("::")
        if separator_index == -1:
            value = payload
        else:
            name = payload[:separator_index]
            value = payload[separator_index + 2:]

    if not value:
        return None
    return {"nome": name or "Arquivo", "url": value}

def _drive_url_from_file(file_info: dict) -> str:
    url = (file_info.get("url") or "").strip()
    if url.startswith("http://") or url.startswith("https://"):
        return url
    fid = file_info.get("drive_file_id")
    return f"https://drive.google.com/file/d/{fid}/view" if fid else ""

_ACTION_SEARCH_STOPWORDS = {
    "de", "a", "o", "que", "e", "do", "da", "em", "um", "uma",
    "os", "as", "no", "na", "com", "por", "para", "dos", "das",
    "nos", "nas", "ao", "se", "ou", "acao", "acoes", "tarefa", "tarefas",
    "pesquisa", "pesquisar", "pesquise", "busca", "buscar", "busque",
    "procura", "procurar", "procure", "localiza", "localizar", "localize",
    "ative", "ativar", "ativa", "contexto", "hermes", "por", "favor",
}

def _action_query_terms(text: str) -> list[str]:
    normalized = _normalize_for_matching(text)
    return [
        term
        for term in re.findall(r"[\w./-]+", normalized, flags=re.UNICODE)
        if term not in _ACTION_SEARCH_STOPWORDS and len(term) > 2
    ]

def _clean_action_search_query(text: str) -> str:
    query = _normalize_for_matching(text)
    query = re.sub(r"\b(?:hermes|por favor)\b", " ", query, flags=re.IGNORECASE)
    query = re.sub(r"\b(?:pesquise|pesquisa|pesquisar|busque|busca|buscar|procure|procura|procurar|localize|localiza|localizar)\b", " ", query, flags=re.IGNORECASE)
    query = re.sub(r"\b(?:a|uma|o|um)?\s*(?:acao|acoes|tarefa|tarefas)\b", " ", query, flags=re.IGNORECASE)
    query = re.sub(r"\b(?:e\s+)?(?:ative|ativa|ativar|entre|entra|entrar)\b.*\bcontexto\b", " ", query, flags=re.IGNORECASE)
    query = re.sub(r"\b(?:no|na|em|o)?\s*contexto\b", " ", query, flags=re.IGNORECASE)
    query = re.sub(r"\s+", " ", query)
    return query.strip(" .,:;-")

def _extract_action_search_context_query(text: str) -> str | None:
    lowered = _normalize_for_matching(text).strip()
    if not lowered:
        return None
    # Long messages (forwarded texts, context info) are never intent commands
    if len(text.strip()) > 500:
        return None
    wants_search = any(re.search(r'\b' + marker + r'\b', lowered) for marker in (
        "pesquisa", "pesquise", "buscar", "busca", "busque",
        "procura", "procure", "localiza", "localize",
    ))
    mentions_action = any(marker in lowered for marker in ("acao", "acoes", "tarefa", "tarefas"))
    wants_context = "contexto" in lowered and any(re.search(r'\b' + marker + r'\b', lowered) for marker in (
        "ative", "ativa", "ativar", "entre", "entra", "entrar",
    ))
    if not (wants_search and mentions_action and wants_context):
        return None
    cleaned = _clean_action_search_query(text)
    return cleaned or None

def _extract_action_lookup_query(text: str) -> str | None:
    """Detects requests that only ask to find/list an action, without locking context."""
    lowered = _normalize_for_matching(text).strip()
    if not lowered:
        return None
    if len(text.strip()) > 500:
        return None
    wants_search = any(marker in lowered for marker in (
        "pesquisa", "pesquise", "pesquisar", "buscar", "busca", "busque",
        "procura", "procurar", "procure", "localiza", "localizar", "localize",
        "encontra", "encontrar", "encontre", "ache", "achar",
    ))
    mentions_action = any(marker in lowered for marker in ("acao", "acoes", "tarefa", "tarefas"))
    wants_context = "contexto" in lowered and any(marker in lowered for marker in (
        "ative", "ativa", "ativar", "entre", "entra", "entrar",
    ))
    if not (wants_search and mentions_action) or wants_context:
        return None
    cleaned = _clean_action_search_query(text)
    cleaned = re.sub(r"\b(?:relacionad[ao]s?|referente|sobre|chamad[ao]s?|nomead[ao]s?)\b", " ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .,:;-")
    return cleaned or None

def _extract_natural_context_query(text: str) -> str | None:
    """Detects natural-language requests such as 'entre no contexto da acao X'."""
    lowered = _normalize_for_matching(text).strip()
    if not lowered:
        return None

    patterns = [
        r"^(?:entre|entra|entrar)\s+(?:no|na|em)?\s*contexto\s*(?:da|do|de)?\s*(?:acao|tarefa)?\s*(?:da|do|de)?\s*(.*)$",
        r"^(?:ative|ativa|ativar)\s+(?:o\s+)?contexto\s*(?:da|do|de)?\s*(?:acao|tarefa)?\s*(?:da|do|de)?\s*(.*)$",
    ]
    for pattern in patterns:
        match = re.match(pattern, lowered, flags=re.IGNORECASE)
        if match:
            return (match.group(1) or "").strip(" .,:;-")
    return None

def _is_list_actions_request(text: str) -> tuple[bool, str | None]:
    """Detects 'liste as ações de hoje' style requests.
    Returns (is_match, date_iso_or_None) — date_iso is None when no date was specified.
    """
    lowered = _normalize_for_matching(text).strip()
    if not lowered or len(text.strip()) > 250:
        return False, None

    list_markers = (
        "liste", "listar", "me liste",
        "mostre", "mostrar", "me mostre", "mostra", "me mostra",
        "me da", "me de",
        "quais sao",
    )
    action_markers = ("acoes", "tarefas", "atividades")

    has_list = any(marker in lowered for marker in list_markers)
    has_action = any(marker in lowered for marker in action_markers)

    wants_context = "contexto" in lowered and any(m in lowered for m in ("ative", "ativa", "entre", "entra"))
    has_file = any(m in lowered for m in ("arquivo", "arquivos", "documento", "documentos", "link", "links", "anexo", "anexos"))

    if not (has_list and has_action) or wants_context or has_file:
        return False, None

    date_iso = _requested_date_from_text(text)
    return True, date_iso

def _fetch_actions_by_date(date_iso: str | None) -> list:
    """Fetches actions filtered by data_limite == date_iso, or all if date_iso is None."""
    from tools.busca_grafo import buscar_tarefas
    if date_iso:
        res = buscar_tarefas(
            "",
            match_mode="any",
            limite=200,
            data_limite_inicio=date_iso,
            data_limite_fim=date_iso,
        )
    else:
        res = buscar_tarefas("", match_mode="any", limite=200)
    return res.get("resultados", [])

def _format_actions_simple_list(results: list, date_iso: str | None) -> str:
    """Formats a plain bullet-point list of action titles."""
    if date_iso:
        try:
            d = datetime.fromisoformat(date_iso)
            date_label = d.strftime("%d/%m/%Y")
        except Exception:
            date_label = date_iso
        header = f"Ações para <b>{date_label}</b>:"
    else:
        header = "Ações cadastradas:"

    if not results:
        return header + "\n\n<i>Nenhuma ação encontrada.</i>"

    lines = [header, ""]
    for r in results:
        titulo = html.escape(r.get("titulo") or "sem título")
        lines.append(f"• {titulo}")

    return "\n".join(lines)

def _is_reset_request(text: str) -> bool:
    """Detecta se o usuário quer resetar/limpar a sessão via linguagem natural."""
    lowered = _normalize_for_matching(text).strip()
    if not lowered:
        return False
    
    # Marcadores de ação
    reset_markers = ["limpar", "limpa", "resetar", "reset", "recomecar", "recomeca", "reiniciar", "reinicia"]
    # Marcadores de objeto
    context_markers = ["chat", "conversa", "historico", "sessao", "memoria"]
    
    has_reset = any(m in lowered for m in reset_markers)
    has_context = any(m in lowered for m in context_markers)
    
    # Se disser apenas "reset" ou "limpa tudo", ou combinar reset + contexto
    return (has_reset and has_context) or (has_reset and len(lowered.split()) <= 2)

def _select_context_result(query: str, results: list) -> dict | None:
    if not results:
        return None
    if len(results) == 1:
        return results[0]

    norm_query = _normalize_for_matching(query).strip()
    if not norm_query:
        return None

    exact = [r for r in results if _normalize_for_matching(r.get("titulo", "")) == norm_query]
    if len(exact) == 1:
        return exact[0]

    contained = [r for r in results if norm_query in _normalize_for_matching(r.get("titulo", ""))]
    if len(contained) == 1:
        return contained[0]

    query_terms = set(_action_query_terms(query))
    if query_terms:
        scored = []
        for r in results:
            title_terms = set(_action_query_terms(r.get("titulo", "")))
            overlap = len(query_terms & title_terms)
            scored.append((overlap, r))
        scored.sort(key=lambda item: item[0], reverse=True)
        best_score = scored[0][0]
        second_score = scored[1][0] if len(scored) > 1 else 0
        if best_score >= 2 and best_score > second_score:
            return scored[0][1]
    return None

def _lock_action_session(session: dict, task_id: str, snapshot: dict) -> None:
    session["contexto_ativo"] = "acao"
    session["acao_id"] = task_id
    session["acao_titulo"] = snapshot.get("titulo") or task_id
    session["acao_context_snapshot"] = snapshot
    session["history_acao"] = []

def _format_context_locked_message(snapshot: dict) -> str:
    titulo = html.escape(snapshot.get("titulo") or "acao")
    arquivos_count = len(snapshot.get("arquivos_disponiveis") or [])
    return (
        f"🔒 <b>[Contexto: {titulo}]</b>\n\n"
        "Contexto trancado. Estou focado exclusivamente nesta ação.\n"
        f"Arquivos detectados no contexto: <b>{arquivos_count}</b>.\n\n"
        "Use <i>Sair do Contexto</i> para retornar ao modo geral."
    )

def _is_context_file_request(text: str) -> bool:
    lowered = _normalize_for_matching(text)
    file_markers = (
        "link",
        "links",
        "documento",
        "documentos",
        "arquivo",
        "arquivos",
        "anexo",
        "anexos",
    )
    request_markers = (
        "qual",
        "quais",
        "mande",
        "manda",
        "envie",
        "enviar",
        "me manda",
        "pode me mandar",
        "listar",
        "liste",
        "mostre",
        "mostrar",
    )
    return any(marker in lowered for marker in file_markers) and any(marker in lowered for marker in request_markers)

def _requested_date_from_text(text: str) -> str | None:
    lowered = _normalize_for_matching(text)
    try:
        from zoneinfo import ZoneInfo
        today = datetime.now(ZoneInfo("America/Sao_Paulo")).date()
    except Exception:
        today = datetime.now(timezone.utc).date()

    if "ontem" in lowered:
        return (today - timedelta(days=1)).isoformat()
    if "hoje" in lowered:
        return today.isoformat()

    match = re.search(r"\b(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\b", lowered)
    if not match:
        return None
    day = int(match.group(1))
    month = int(match.group(2))
    year_raw = match.group(3)
    year = int(year_raw) if year_raw else today.year
    if year < 100:
        year += 2000
    try:
        return datetime(year, month, day).date().isoformat()
    except ValueError:
        return None

def _format_context_files_response(acao_snapshot: dict, request_text: str) -> str:
    arquivos = list(acao_snapshot.get("arquivos_disponiveis") or [])
    date_filter = _requested_date_from_text(request_text)
    if date_filter:
        arquivos = [
            a for a in arquivos
            if str(a.get("data_diario") or a.get("data_criacao") or "").startswith(date_filter)
        ]

    lowered = _normalize_for_matching(request_text)
    requested_terms = []
    if "deferid" in lowered:
        requested_terms.append("deferid")
    if "indeferid" in lowered:
        requested_terms.append("indeferid")
    if requested_terms:
        term_filtered = [
            a for a in arquivos
            if any(term in _normalize_for_matching(a.get("nome", "")) for term in requested_terms)
        ]
        if term_filtered:
            arquivos = term_filtered

    titulo = html.escape(acao_snapshot.get("titulo") or "acao")
    if not arquivos:
        when = f" em {date_filter}" if date_filter else ""
        return (
            f"<b>[Contexto: {titulo}]</b>\n\n"
            f"Nao encontrei anexos ou links registrados{when} nesse contexto.\n"
            "Nao vou criar links de exemplo. Se o arquivo estiver visivel no diario mas nao aparecer aqui, "
            "provavelmente ele nao foi gravado no pool/snapshot da tarefa."
        )

    lines = [
        f"<b>[Contexto: {titulo}]</b>",
        "",
        "Encontrei estes arquivos registrados no contexto. Estou mostrando a URL completa para evitar link oculto incorreto:",
        "",
    ]
    for file_info in arquivos[:12]:
        name = html.escape(file_info.get("nome") or "Arquivo sem nome")
        url = html.escape(_drive_url_from_file(file_info), quote=False)
        date_label = (file_info.get("data_diario") or file_info.get("data_criacao") or "")[:10]
        date_suffix = f" ({html.escape(date_label)})" if date_label else ""
        if url:
            lines.append(f"- <b>{name}</b>{date_suffix}\n  {url}")
        else:
            fid = html.escape(file_info.get("drive_file_id") or "sem drive_file_id")
            lines.append(f"- <b>{name}</b>{date_suffix}\n  DRIVE_FILE_ID: <code>{fid}</code>")

    if len(arquivos) > 12:
        lines.append(f"\nMais {len(arquivos) - 12} arquivo(s) omitido(s) para caber no Telegram.")
    return "\n".join(lines)

def _fetch_acao_snapshot(db, task_id: str) -> dict | None:
    """Busca e formata dados de uma tarefa para persistir no snapshot da sessão."""
    try:
        from main import get_drive_service

        doc = db.collection("tarefas").document(task_id).get()
        if not doc.exists:
            return None
        data = doc.to_dict() or {}

        acomp_raw = data.get("acompanhamento") or []
        acomp_entries = []
        diario_full = []

        def _obter_data_ordenacao(entry):
            d = entry.get('data')
            if not d:
                return ""
            if isinstance(d, str):
                return d
            import datetime as _datetime
            if isinstance(d, _datetime.datetime):
                if d.tzinfo is None:
                    d = d.replace(tzinfo=_datetime.timezone.utc)
                return d.isoformat()
            if isinstance(d, _datetime.date):
                return d.isoformat()
            if hasattr(d, 'isoformat'):
                try:
                    return d.isoformat()
                except Exception:
                    pass
            return str(d)

        def _formatar_data_curta(entry):
            d_val = entry.get('data')
            if not d_val:
                return ""
            if isinstance(d_val, str):
                return d_val[:10]
            elif hasattr(d_val, "strftime"):
                try:
                    return d_val.strftime("%Y-%m-%d")
                except Exception:
                    return str(d_val)[:10]
            else:
                return str(d_val)[:10]

        for entry in sorted(acomp_raw, key=_obter_data_ordenacao):
            nota = (entry.get("nota") or "").strip()
            data_entry = _formatar_data_curta(entry)
            if nota:
                diario_full.append(f"[{data_entry}] {nota[:500]}")
        for entry in sorted(acomp_raw, key=_obter_data_ordenacao, reverse=True)[:3]:
            nota = (entry.get("nota") or "").strip()
            data_entry = _formatar_data_curta(entry)
            if nota:
                acomp_entries.append(f"[{data_entry}] {nota[:300]}")

        plano_raw = data.get("plano_acao") or []
        plano_entries = []
        for passo in plano_raw:
            if isinstance(passo, dict):
                texto = (passo.get("text") or passo.get("titulo") or "").strip()
                concluido = passo.get("completed", False)
                if texto:
                    marcador = "✓" if concluido else "○"
                    plano_entries.append(f"{marcador} {texto[:200]}")
            elif isinstance(passo, str) and passo.strip():
                plano_entries.append(f"○ {passo.strip()[:200]}")

        tags = data.get("kg_tags") or data.get("tags") or []
        arquivos_disponiveis = []
        arquivos_seen = set()
        drive_service = None

        def add_file_candidate(nome: str, url: str, data_criacao: str = "", data_diario: str = ""):
            nonlocal drive_service
            fid = _extract_drive_file_id(url)
            if not fid:
                return
            key = fid or f"{nome}|{url}"
            if key in arquivos_seen:
                for existing in arquivos_disponiveis:
                    if existing.get("drive_file_id") == fid:
                        if data_diario and not existing.get("data_diario"):
                            existing["data_diario"] = data_diario
                        if data_criacao and not existing.get("data_criacao"):
                            existing["data_criacao"] = data_criacao
                        if url and not existing.get("url"):
                            existing["url"] = url
                        break
                return
            if drive_service is None:
                try:
                    drive_service = get_drive_service()
                except Exception:
                    drive_service = False
            if drive_service:
                try:
                    meta = drive_service.files().get(fileId=fid, fields="id,trashed").execute()
                    if meta.get("trashed"):
                        return
                except Exception:
                    return
            arquivos_seen.add(key)
            arquivos_disponiveis.append({
                "nome": nome or "Arquivo sem nome",
                "drive_file_id": fid,
                "url": url,
                "data_criacao": data_criacao or "",
                "data_diario": data_diario or "",
            })

        for item in data.get("pool_dados", []):
            if item.get("tipo") != "arquivo":
                continue
            fid = item.get("drive_file_id")
            raw_url = item.get("valor", "") or ""
            url = raw_url if _extract_drive_file_id(raw_url) else (f"https://drive.google.com/file/d/{fid}/view" if fid else raw_url)
            add_file_candidate(
                item.get("nome", "Arquivo sem nome"),
                url,
                data_criacao=str(item.get("data_criacao") or ""),
            )

        for entry in acomp_raw:
            parsed_file = _parse_diary_file_note(entry.get("nota"))
            if not parsed_file:
                continue
            add_file_candidate(
                parsed_file.get("nome") or "Arquivo",
                parsed_file.get("url") or "",
                data_diario=str(entry.get("data") or ""),
            )

        return {
            "id": task_id,
            "titulo": data.get("titulo", "sem titulo"),
            "status": data.get("status", ""),
            "area": data.get("area_tematica", ""),
            "area_tematica": data.get("area_tematica", ""),
            "descricao": (data.get("descricao") or "")[:500],
            "notas": (data.get("notas") or "")[:400],
            "sintese": (data.get("sintese_demanda") or data.get("demanda") or "")[:400],
            "plano_atual": data.get("plano_acao", []),
            "plano_acao": plano_entries,
            "diario_integral": "\n".join(diario_full)[-4000:],
            "acompanhamento_recente": acomp_entries,
            "tags": tags,
            "arquivos_disponiveis": arquivos_disponiveis[:20],
        }
    except Exception as e:
        print(f"[Session] Falha ao buscar tarefa {task_id}: {e}")
        return None

def _cached_acao_snapshot(db, task_id: str) -> dict | None:
    now = time.monotonic()
    cached = _ACTION_SNAPSHOT_CACHE.get(task_id)
    if cached and (now - cached[0]) < _ACTION_SNAPSHOT_CACHE_TTL:
        return cached[1]

    snapshot = _fetch_acao_snapshot(db, task_id)
    if snapshot:
        _ACTION_SNAPSHOT_CACHE[task_id] = (now, snapshot)
    return snapshot

def _search_actions_for_context(db, query: str) -> list:
    """Busca tarefas usando buscar_tarefas e retorna até 5 resultados para seleção."""
    from tools.busca_grafo import buscar_tarefas

    if query:
        terms = _action_query_terms(query)
        mode = "all" if len(terms) >= 2 else "any"
        res = buscar_tarefas(query, match_mode=mode, limite=5)
        if mode == "all" and not res.get("resultados"):
            res = buscar_tarefas(query, match_mode="any", limite=5)
    else:
        res = buscar_tarefas("", match_mode="any", limite=5)

    return res.get("resultados", [])

def _build_action_keyboard(results: list) -> list:
    """Constrói InlineKeyboard com uma linha por resultado de ação."""
    keyboard = []
    for r in results:
        task_id = r["id"]
        titulo = (r.get("titulo") or "sem titulo")[:35]
        area = r.get("area", "")
        label = titulo + (f" [{area}]" if area else "")
        keyboard.append([{"text": label[:60], "callback_data": f"lock:{task_id}"}])
    return keyboard

def _format_action_lookup_results(query: str, results: list) -> str:
    query_label = html.escape(query or "sua busca")
    if not results:
        return (
            f"Nenhuma acao encontrada para <i>{query_label}</i> no historico do Hermes.\n"
            "Nao pesquisei e-mails, acervo ou internet porque voce pediu uma acao do sistema."
        )

    if len(results) == 1:
        header = f"Encontrei esta acao para <i>{query_label}</i>:"
    else:
        header = f"Encontrei {len(results)} acoes mais provaveis para <i>{query_label}</i>:"

    lines = [header]
    for idx, r in enumerate(results[:5], 1):
        titulo = html.escape(r.get("titulo") or "sem titulo")
        task_id = html.escape(str(r.get("id") or ""))
        status = html.escape(r.get("status") or "nao informado")
        area = html.escape(r.get("area") or "nao informada")
        prazo = html.escape(str(r.get("data_limite") or "nao informado"))
        lines.append(f"\n<b>{idx}. {titulo}</b>")
        lines.append(f"ID: <code>{task_id}</code>")
        lines.append(f"Status: {status} | Area: {area} | Prazo: {prazo}")
        if r.get("processo_sei"):
            lines.append(f"Processo SEI: {html.escape(str(r.get('processo_sei')))}")
        lines.append(f"Abrir: task:{task_id}")
    return "\n".join(lines)

_EXIT_KEYBOARD = [[{"text": "🔓 Sair do Contexto", "callback_data": "exit_context"}]]

_RESET_CONFIRM_KEYBOARD = [
    [
        {"text": "✅ Sim, limpar", "callback_data": "reset_session_confirm"},
        {"text": "❌ Não, manter", "callback_data": "reset_session_cancel"}
    ]
]

_CONFIRM_ACAO_KEYBOARD = [
    [
        {"text": "✅ Confirmar Registro", "callback_data": "confirm_acao"},
        {"text": "❌ Cancelar", "callback_data": "cancel_acao"}
    ]
]

_CONFIRM_FINANCEIRO_KEYBOARD = [
    [
        {"text": "✅ Confirmar Lançamento", "callback_data": "confirm_financeiro"},
        {"text": "❌ Cancelar", "callback_data": "cancel_financeiro"}
    ]
]

_CONFIRM_WHATSAPP_KEYBOARD = [
    [
        {"text": "✅ Confirmar Envio", "callback_data": "confirm_whatsapp"},
        {"text": "❌ Cancelar", "callback_data": "cancel_whatsapp"}
    ]
]

_HEALTH_RADICULAR_LOCATIONS = [
    ("nenhum", "Nenhum"), ("gluteo", "Glúteo"), ("quadril", "Quadril"), ("coxa", "Coxa"),
    ("joelho", "Joelho"), ("panturrilha", "Panturrilha"), ("tornozelo", "Tornozelo"), ("pe", "Pé"),
]

_HEALTH_THERAPY_MODALITIES = [
    ("pilates", "Pilates"), ("fisioterapia", "Fisioterapia"), ("rpg", "RPG"),
    ("acupuntura", "Acupuntura"), ("nenhuma", "Nenhuma"),
]

_HEALTH_TRIGGER_TYPES = [
    ("espirro_crise_alergica", "Espirro/crise alérgica"), ("viagem_longa_sentado", "Viagem longa sentado"),
    ("dia_muito_sentado", "Dia muito sentado"), ("torcao_no_sono", "Torção no sono"),
    ("carga_assimetrica", "Carga assimétrica"), ("estresse", "Estresse"), ("outro", "Outro"),
    ("nenhum", "Nenhum"),
]

_HEALTH_RED_FLAG_MESSAGE = (
    "⚠️ <b>Sinais de alerta no registro de hoje</b>\n\n"
    "Isto não é um diagnóstico — é uma lista de sinais que, se presentes, indicam buscar "
    "atendimento médico o quanto antes:\n"
    "• Dormência na região da sela (períneo)\n"
    "• Alteração no controle de urina ou fezes\n"
    "• Fraqueza progressiva para levantar a ponta do pé\n"
    "• Sintomas nas duas pernas ao mesmo tempo\n\n"
    "Se notar qualquer um desses, procure atendimento."
)

def _health_entry_source_after_telegram(current_data: dict) -> str:
    """'ambos' se o painel ja tocou o registro do dia; senao 'telegram'."""
    return "ambos" if current_data.get("entrySource") == "painel" else "telegram"

def _transcribe_audio_bytes(audio_bytes: bytes, extension: str, db) -> str:
    """Transcreve áudio usando Groq (Whisper) com fallback para Gemini."""
    import base64 as _b64
    import tempfile
    import os
    keys = _get_api_keys(db)
    
    # 1. Tentativa com Groq (Whisper-Large-V3-Turbo) - Alta qualidade para PT-BR
    groq_key = keys.get("groq_api_key")
    if groq_key:
        try:
            from groq import Groq
            client = Groq(api_key=groq_key)
            suffix = extension if extension.startswith(".") else f".{extension}"
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(audio_bytes)
                tmp_path = tmp.name
            try:
                with open(tmp_path, "rb") as f:
                    transcription = client.audio.transcriptions.create(
                        file=(os.path.basename(tmp_path), f),
                        model="whisper-large-v3-turbo",
                        response_format="json",
                        language="pt",
                        temperature=0.0,
                    )
                if transcription and hasattr(transcription, "text") and transcription.text:
                    return transcription.text.strip()
            finally:
                try:
                    if os.path.exists(tmp_path):
                        os.remove(tmp_path)
                except OSError:
                    pass
        except Exception as groq_err:
            print(f"[Transcription] Groq failed: {groq_err}")

    # 2. Fallback: Gemini (Multimodal)
    gemini_key = keys.get("gemini_api_key")
    if gemini_key:
        try:
            from google import genai
            from google.genai import types
            client = genai.Client(api_key=gemini_key)
            
            # Mapeamento básico de extensões para mime-types
            mime_map = {
                "ogg": "audio/ogg", "oga": "audio/ogg", "wav": "audio/wav",
                "mp3": "audio/mpeg", "m4a": "audio/mp4", "flac": "audio/flac"
            }
            clean_ext = extension.lower().strip(".")
            mime_type = mime_map.get(clean_ext, "audio/ogg")
            
            response = generate_content_logged(
                client,
                model=GEMINI_TRANSCRIPTION_MODEL,
                contents=[
                    types.Part.from_bytes(data=audio_bytes, mime_type=mime_type),
                    "Transcreva este áudio literalmente. Responda apenas com o texto transcrito, sem introduções ou explicações."
                ],
                feature="telegram.audio_transcription_fallback",
            )
            if response and response.text:
                return response.text.strip()
        except Exception as gemini_err:
            print(f"[Transcription] Gemini fallback failed: {gemini_err}")

    return "[Transcrição indisponível: erro nos motores Groq/Gemini]"

_VIDEO_INLINE_MAX_BYTES = 10 * 1024 * 1024  # acima disso usa Files API (teto ~20MB de payload inline do Gemini, já com folga p/ inflação do base64)

_VIDEO_MIME_MAP = {
    # cobre extensões convencionais E os subtipos de MIME crus que o worker
    # de captura grava (mimetype.split('/')[1] — ex.: video/quicktime -> .quicktime)
    "mp4": "video/mp4", "quicktime": "video/quicktime", "mov": "video/quicktime",
    "webm": "video/webm", "3gpp": "video/3gpp", "3gp": "video/3gpp",
    "mpeg": "video/mpeg", "mpg": "video/mpeg",
}

_VIDEO_TRANSCRIPTION_PROMPT = (
    "Transcreva a fala deste vídeo literalmente e, em seguida, descreva em 1 a 3 frases "
    "o conteúdo visual relevante (documentos, produtos, telas gravadas, ações). "
    "Responda em português, direto, sem introduções."
)

def _transcribe_video_bytes(video_bytes: bytes, extension: str, db) -> str:
    """Transcreve fala + descreve conteúdo visual de um vídeo via Gemini nativo
    (entendimento multimodal direto — sem extrair áudio via FFmpeg)."""
    keys = _get_api_keys(db)
    gemini_key = keys.get("gemini_api_key")
    if not gemini_key:
        return "[Transcrição indisponível: Gemini API Key não configurada]"

    clean_ext = extension.lower().strip(".")
    mime_type = _VIDEO_MIME_MAP.get(clean_ext, "video/mp4")

    try:
        from google import genai
        from google.genai import types
        client = genai.Client(api_key=gemini_key)

        if len(video_bytes) <= _VIDEO_INLINE_MAX_BYTES:
            response = generate_content_logged(
                client,
                model=GEMINI_TRANSCRIPTION_MODEL,
                contents=[types.Part.from_bytes(data=video_bytes, mime_type=mime_type), _VIDEO_TRANSCRIPTION_PROMPT],
                feature="whatsapp_consolidation.video_transcription",
                db=db,
            )
            if response and response.text:
                return response.text.strip()
            return "[Transcrição indisponível: resposta vazia do Gemini]"

        suffix = f".{clean_ext or 'mp4'}"
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(video_bytes)
            tmp_path = tmp.name
        gemini_file = None
        try:
            gemini_file = client.files.upload(
                file=tmp_path,
                config=types.UploadFileConfig(mime_type=mime_type, display_name=os.path.basename(tmp_path)),
            )
            waited = 0
            while str(getattr(gemini_file.state, "name", gemini_file.state)) == "PROCESSING":
                if waited >= 120:
                    raise TimeoutError("Files API demorou demais para preparar o vídeo (>2min).")
                time.sleep(3)
                waited += 3
                gemini_file = client.files.get(name=gemini_file.name)
            final_state = str(getattr(gemini_file.state, "name", gemini_file.state))
            if final_state == "FAILED":
                raise RuntimeError("Files API falhou ao preparar o vídeo.")

            response = generate_content_logged(
                client,
                model=GEMINI_TRANSCRIPTION_MODEL,
                contents=[_VIDEO_TRANSCRIPTION_PROMPT, gemini_file],
                feature="whatsapp_consolidation.video_transcription",
                db=db,
            )
            if response and response.text:
                return response.text.strip()
            return "[Transcrição indisponível: resposta vazia do Gemini]"
        finally:
            if gemini_file is not None:
                try:
                    client.files.delete(name=gemini_file.name)
                except Exception:
                    pass
            try:
                if os.path.exists(tmp_path):
                    os.remove(tmp_path)
            except OSError:
                pass
    except Exception as exc:
        print(f"[Transcription] Gemini video understanding failed: {exc}")

    return "[Transcrição indisponível: erro no motor Gemini]"

def _blob_public_url(blob) -> str:
    """
    Returns a public URL for a Firebase Storage blob.
    Generates and patches a Firebase Storage download token metadata to ensure the link is publicly accessible.
    """
    import urllib.parse
    import uuid
    
    token = None
    if blob.metadata:
        token = blob.metadata.get("firebaseStorageDownloadTokens")
        
    if not token:
        token = str(uuid.uuid4())
        metadata = dict(blob.metadata or {})
        metadata["firebaseStorageDownloadTokens"] = token
        blob.metadata = metadata
        try:
            blob.patch()
        except Exception as e:
            print(f"[Storage] Warning: Failed to patch blob metadata for token: {e}")
            
    bucket_name = blob.bucket.name
    encoded_name = urllib.parse.quote(blob.name, safe='')
    if token:
        return f"https://firebasestorage.googleapis.com/v0/b/{bucket_name}/o/{encoded_name}?alt=media&token={token}"
    return f"https://firebasestorage.googleapis.com/v0/b/{bucket_name}/o/{encoded_name}?alt=media"

def _upload_to_storage(file_bytes: bytes, file_name: str, mime_type: str, chat_id: str) -> str:
    """Uploads to Firebase Storage and returns the public URL."""
    bucket = _get_hermes_storage_bucket()
    path = f"telegram_uploads/{chat_id}/{file_name}"
    blob = bucket.blob(path)
    blob.upload_from_string(file_bytes, content_type=mime_type)
    return _blob_public_url(blob)

def _store_telegram_media(file_bytes: bytes, file_name: str, mime_type: str, chat_id: str) -> tuple[str | None, str | None]:
    """
    Stores media inline when small enough; otherwise uploads to Storage and returns the blob path.
    Returns: (media_bytes_b64, storage_path)
    """
    import base64

    if len(file_bytes) <= _MAX_INLINE_MEDIA_BYTES:
        return base64.b64encode(file_bytes).decode(), None

    safe_name = os.path.basename(file_name or f"upload_{int(time.time())}")
    bucket = _get_hermes_storage_bucket()
    path = f"telegram_uploads/{chat_id}/{int(time.time())}_{safe_name}"
    blob = bucket.blob(path)
    blob.upload_from_string(file_bytes, content_type=mime_type)
    return None, path

def _load_telegram_media_bytes(media_bytes_b64: str | None, storage_path: str | None) -> bytes | None:
    if media_bytes_b64:
        import base64

        return base64.b64decode(media_bytes_b64)
    if storage_path:
        bucket = _get_hermes_storage_bucket()
        blob = bucket.blob(storage_path)
        return blob.download_as_bytes()
    return None

_AREAS_VALIDAS_CACHE: tuple | None = None  # (monotonic_ts, list[str])

_AREAS_VALIDAS_TTL = 300  # unidades mudam raramente; evita stream da coleção a cada mensagem

def carregar_areas_tematicas_validas(db) -> list[str]:
    """Lista canonica de areas tematicas = nomes das Unidades + GERAL/NAO CLASSIFICADA."""
    global _AREAS_VALIDAS_CACHE
    now = time.monotonic()
    if _AREAS_VALIDAS_CACHE and (now - _AREAS_VALIDAS_CACHE[0]) < _AREAS_VALIDAS_TTL:
        return _AREAS_VALIDAS_CACHE[1]
    nomes = []
    try:
        for d in db.collection('unidades').stream():
            nome = ((d.to_dict() or {}).get('nome') or '').strip().upper()
            if nome:
                nomes.append(nome)
    except Exception as e:
        print(f"[areas_validas] Falha ao carregar unidades: {e}")
    base = ['GERAL', 'NÃO CLASSIFICADA']
    # preserva ordem e remove duplicados
    resultado = base + [n for n in nomes if n not in base]
    if nomes:  # não cacheia resultado de falha
        _AREAS_VALIDAS_CACHE = (now, resultado)
    return resultado

def normalizar_area_tematica(valor: str, areas_validas: list[str]) -> str:
    """Casa case-insensitive contra a lista valida; fallback 'GERAL'."""
    alvo = (valor or '').strip().upper()
    for a in areas_validas:
        if a.upper() == alvo:
            return a
    return 'GERAL'

def _build_system_instruction(copilot_core: str, copilot_soul: str, contexto_ativo: str) -> str:
    from datetime import datetime as _dt
    from zoneinfo import ZoneInfo
    tz = ZoneInfo("America/Sao_Paulo")
    now_local = _dt.now(tz)
    today = now_local.strftime("%Y-%m-%d")
    time_str = now_local.strftime("%H:%M")
    ctx_hint = (
        f"\n\n<b>Contexto ativo:</b> {contexto_ativo}"
        if contexto_ativo != "geral"
        else ""
    )
    return (
        f"Você é o Copiloto Hermes, estrategista sênior de processos. Hoje é {today} e o horário local atual é {time_str}."
        f"{ctx_hint}\n\n"
        "## CORE ESTÁTICO DO COPILOTO\n"
        f"{copilot_core}\n\n"
        "## PERSONALIDADE DINÂMICA ATUAL\n"
        f"{copilot_soul}\n\n"
        "## POSTURA CRÍTICA (SOCRÁTICA)\n"
        "Não seja condescendente. Antes de validar, cheque premissas e aponte pontos frágeis, riscos e o que falta — "
        "de forma curta. Quando a decisão tiver peso, faça no máximo UMA pergunta socrática afiada. Não bajule por "
        "reflexo; discorde com fundamento quando os fatos não sustentarem o pedido. Calibre a criticidade ao risco.\n\n"
        "## CANAL DE COMUNICAÇÃO: TELEGRAM\n"
        "REGRA ABSOLUTA DE FORMATAÇÃO: Você está respondendo via Telegram. "
        "Use EXCLUSIVAMENTE as seguintes tags HTML suportadas pelo Telegram: "
        "<b>negrito</b>, <i>itálico</i>, <code>código inline</code>, <pre>bloco de código</pre>. "
        "PROIBIDO usar Markdown (asteriscos, underlines, backticks, #, ##). "
        "Listas: use • ou - como prefixo de linha, sem Markdown. "
        "Mantenha respostas concisas — Telegram tem limite de 4096 caracteres por mensagem.\n\n"
        "## REGRAS ABSOLUTAS\n"
        "1. JAMAIS expanda siglas arbitrariamente.\n"
        "2. Se qualquer ferramenta retornar campo 'erro', reproduza o erro literal.\n"
        "3. Agendamento/Agenda: Por padrão, ao propor ou criar uma ação, NÃO proponha nem defina horários de início/fim (campos horario_inicio e horario_fim devem ser nulos/vazios), definindo apenas o dia (data_limite). SÓ preencha horario_inicio e horario_fim se o usuário pedir explicitamente para agendar um horário específico. Nesse caso, você DEVE usar consultar_agenda ou encontrar_slot_livre ANTES de agendar. Horário de funcionamento: 08:00 às 19:00, janela D+7. Se o agendamento for para hoje, o horário inicial DEVE ser sempre posterior ao horário local atual. Se houver conflito em horário específico, pergunte se força inserção ou busca outro slot.\n"
        "4. Para criar uma nova ação, você DEVE usar obrigatoriamente a ferramenta propor_acao_para_confirmacao. Ela gerará botões de ✅/❌ para o usuário confirmar o registro.\n"
        "5. Links de tarefas: use o formato task:{ID} no texto (ex: 'Ação task:abc123').\n"
        "6. Acione salvar_memoria_global apenas para fatos duráveis e preferências estáveis.\n"
        "7. ACESSO FINANCEIRO: para consultas sobre balanço, rendas, obrigações ou metas, use consultar_financas_v2. Para novos registros, use obrigatoriamente propor_lancamento_financeiro para que o usuário receba botões de confirmação. NUNCA invente números.\n"
    )

def _build_system_instruction_guarded(
    copilot_core: str,
    copilot_soul: str,
    contexto_ativo: str,
    acao_snapshot: dict | None = None,
) -> str:
    from datetime import datetime as _dt
    from zoneinfo import ZoneInfo
    tz = ZoneInfo("America/Sao_Paulo")
    now_local = _dt.now(tz)
    today = now_local.strftime("%Y-%m-%d")
    time_str = now_local.strftime("%H:%M")
    ctx_hint = (
        f"\n\n<b>Contexto ativo:</b> {contexto_ativo}"
        if contexto_ativo != "geral"
        else ""
    )
    base = (
        f"Voce e o Copiloto Hermes, estrategista senior de processos. Hoje e {today} e o horario local atual e {time_str}."
        f"{ctx_hint}\n\n"
        "## CORE ESTATICO DO COPILOTO\n"
        f"{copilot_core}\n\n"
        "## PERSONALIDADE DINAMICA ATUAL\n"
        f"{copilot_soul}\n\n"
        "## POSTURA CRITICA (SOCRATICA)\n"
        "Nao seja condescendente. Antes de validar, cheque premissas e aponte pontos frageis, riscos e o que falta - "
        "de forma curta. Quando a decisao tiver peso, faca no maximo UMA pergunta socratica afiada. Nao bajule por "
        "reflexo; discorde com fundamento quando os fatos nao sustentarem o pedido. Calibre a criticidade ao risco.\n\n"
        "## CANAL DE COMUNICACAO: TELEGRAM\n"
        "REGRA ABSOLUTA DE FORMATACAO: Voce esta respondendo via Telegram. "
        "Use EXCLUSIVAMENTE as seguintes tags HTML suportadas pelo Telegram: "
        "<b>negrito</b>, <i>italico</i>, <code>codigo inline</code>, <pre>bloco de codigo</pre>. "
        "PROIBIDO usar Markdown. "
        "Listas: use • ou - como prefixo de linha, sem Markdown. "
        "Mantenha respostas concisas; Telegram tem limite de 4096 caracteres por mensagem.\n\n"
        "## REGRAS ABSOLUTAS\n"
        "1. JAMAIS expanda siglas arbitrariamente.\n"
        "2. Se qualquer ferramenta retornar campo 'erro', reproduza o erro literal.\n"
        "4. Se o pedido for sobre dados internos do Hermes, tarefas, acoes, historico, agenda ou documentos do sistema, NUNCA use internet como fallback. Nesses casos, use apenas ferramentas internas. Se nao encontrar nada apos usar as ferramentas, admita explicitamente que nao encontrou nos registros do sistema.\n"
        "5. Agendamento/Agenda: Por padrao, ao propor ou criar uma acao, NAO proponha nem defina horarios de inicio/fim (campos horario_inicio e horario_fim devem ser nulos/vazios), definindo apenas o dia (data_limite). SO preencha horario_inicio e horario_fim se o usuario pedir explicitamente para agendar um horario especifico. Nesse caso, voce DEVE usar consultar_agenda ou encontrar_slot_livre ANTES de agendar. Horario de funcionamento: 08:00 as 19:00, janela D+7. Se o agendamento for para hoje, o horario inicial DEVE ser sempre posterior ao horario local atual. Se houver conflito em horario especifico, pergunte se forca insercao ou busca outro slot.\n"
        "6. Para criar uma nova ação, você DEVE usar obrigatoriamente a ferramenta propor_acao_para_confirmacao. Isso apresentará os botões de ✅/❌ ao usuário.\n"
        "7. Links de tarefas: use o formato task:{ID} no texto (ex: 'Acao task:abc123').\n"
        "8. Acione salvar_memoria_global apenas para fatos duraveis e preferencias estaveis.\n"
        "9. GOVERNANCA DE FONTES: ao descrever uma tarefa encontrada por consultar_historico_acoes, use SOMENTE os campos retornados por essa ferramenta. Se um campo nao constar no retorno da ferramenta, diga 'nao informado' em vez de inventar.\n"
        "10. ACESSO FINANCEIRO: use exclusivamente consultar_financas_v2 e para novos registros use propor_lancamento_financeiro para lidar com dados financeiros internos (rendas, contas, metas). Proibido inventar valores.\n"
    )

    if not acao_snapshot:
        return base

    titulo = acao_snapshot.get("titulo", "")
    plano_lines = "\n".join(f"  {p}" for p in acao_snapshot.get("plano_acao", [])) or "  (sem passos registrados)"
    diario_lines = "\n".join(f"  {d}" for d in acao_snapshot.get("acompanhamento_recente", [])) or "  (sem entradas recentes)"
    tags_str = ", ".join(str(t) for t in acao_snapshot.get("tags", [])) or "nenhuma"

    acao_section = (
        "\n\n## ACAO VINCULADA — MODO CONTEXTO TRANCADO\n"
        f"ID: {acao_snapshot.get('id', '')}\n"
        f"Titulo: {titulo}\n"
        f"Area: {acao_snapshot.get('area', '')}\n"
        f"Status: {acao_snapshot.get('status', '')}\n"
        f"Descricao: {acao_snapshot.get('descricao', '') or 'nao informada'}\n"
        f"Sintese: {acao_snapshot.get('sintese', '') or 'nao informada'}\n"
        f"Plano de Acao:\n{plano_lines}\n"
        f"Diario Recente:\n{diario_lines}\n"
        f"Notas: {acao_snapshot.get('notas', '') or 'nenhuma'}\n"
        f"Tags: {tags_str}\n\n"
        "MODO CONTEXTO TRANCADO ATIVADO — REGRAS ABSOLUTAS:\n"
        "1. O historico de conversas anteriores foi isolado. Concentre-se exclusivamente nesta acao.\n"
        "2. Use SOMENTE os dados desta secao ao descrever a acao. PROIBIDO usar RAG ou inferencia.\n"
        "3. Ao receber perguntas sobre a acao, responda com base nos dados acima.\n"
        "4. Se o usuario pedir para atualizar, criar passos ou registrar progresso, use as ferramentas disponiveis.\n"
    )
    return base + acao_section

def _build_system_instruction_guarded_v2(
    copilot_core: str,
    copilot_soul: str,
    contexto_ativo: str,
    acao_snapshot: dict | None = None,
) -> str:
    from datetime import datetime as _dt
    from zoneinfo import ZoneInfo
    tz = ZoneInfo("America/Sao_Paulo")
    now_local = _dt.now(tz)
    today = now_local.strftime("%Y-%m-%d")
    time_str = now_local.strftime("%H:%M")
    ctx_hint = (
        f"\n\n<b>Contexto ativo:</b> {contexto_ativo}"
        if contexto_ativo != "geral"
        else ""
    )
    base = (
        f"Voce e o Copiloto Hermes, estrategista senior de processos. Hoje e {today} e o horario local atual e {time_str}."
        f"{ctx_hint}\n\n"
        "## CORE ESTATICO DO COPILOTO\n"
        f"{copilot_core}\n\n"
        "## PERSONALIDADE DINAMICA ATUAL\n"
        f"{copilot_soul}\n\n"
        "## POSTURA CRITICA (SOCRATICA)\n"
        "Nao seja condescendente. Antes de validar, cheque premissas e aponte pontos frageis, riscos e o que falta - "
        "de forma curta. Quando a decisao tiver peso, faca no maximo UMA pergunta socratica afiada. Nao bajule por "
        "reflexo; discorde com fundamento quando os fatos nao sustentarem o pedido. Calibre a criticidade ao risco.\n\n"
        "## CANAL DE COMUNICACAO: TELEGRAM\n"
        "REGRA ABSOLUTA DE FORMATACAO: Voce esta respondendo via Telegram. "
        "Use EXCLUSIVAMENTE as seguintes tags HTML suportadas pelo Telegram: "
        "<b>negrito</b>, <i>italico</i>, <code>codigo inline</code>, <pre>bloco de codigo</pre>. "
        "PROIBIDO usar Markdown. "
        "Listas: use - ou bullet simples como prefixo de linha, sem Markdown. "
        "Mantenha respostas concisas; Telegram tem limite de 4096 caracteres por mensagem.\n\n"
        "## ESTILO OPERACIONAL E OBJETIVO\n"
        "Responda sem saudacao, preambulo ou encerramento generico. "
        "Comece diretamente pelo dado util, risco, decisao ou proximo passo. "
        "Prefira 1 a 3 bullets curtos quando isso preservar precisao tecnica ou legal. "
        "Nao use frases como 'com base nos dados', 'analisei' ou 'espero ter ajudado'. "
        "Em interacoes proativas, aponte apenas a correlacao, ponto cego ou acao objetiva.\n\n"
        "## REGRAS ABSOLUTAS\n"
        "1. JAMAIS expanda siglas arbitrariamente.\n"
        "2. Se qualquer ferramenta retornar campo 'erro', reproduza o erro literal.\n"
        "4. Se o pedido for sobre dados internos do Hermes, tarefas, acoes, historico, agenda ou documentos do sistema, NUNCA use internet como fallback. Nesses casos, use apenas ferramentas internas. Se nao encontrar nada apos usar as ferramentas, admita explicitamente que nao encontrou nos registros do sistema.\n"
        "5. Agendamento/Agenda: Por padrao, ao propor ou criar uma acao, NAO proponha nem defina horarios de inicio/fim (campos horario_inicio e horario_fim devem ser nulos/vazios), definindo apenas o dia (data_limite). SO preencha horario_inicio e horario_fim se o usuario pedir explicitamente para agendar um horario especifico. Nesse caso, voce DEVE usar consultar_agenda ou encontrar_slot_livre ANTES de agendar. Horario de funcionamento: 08:00 as 19:00, janela D+7. Se o agendamento for para hoje, o horario inicial DEVE ser sempre posterior ao horario local atual. Se houver conflito em horario especifico, pergunte se forca insercao ou busca outro slot.\n"
        "6. Para criar uma nova ação, você DEVE usar obrigatoriamente a ferramenta propor_acao_para_confirmacao. Isso apresentará os botões de ✅/❌ ao usuário.\n"
        "7. Links de tarefas: use o formato task:{ID} no texto (ex: 'Acao task:abc123').\n"
        "8. Acione salvar_memoria_global apenas para fatos duraveis e preferencias estaveis.\n"
        "9. GOVERNANCA DE FONTES: ao descrever uma tarefa encontrada por consultar_historico_acoes, use SOMENTE os campos retornados por essa ferramenta. Se um campo nao constar no retorno da ferramenta, diga 'nao informado' em vez de inventar.\n"
        "10. LINKS E ARQUIVOS: nunca crie hiperlinks, texto-ancora ou URLs que nao aparecam literalmente no contexto, em uma ferramenta ou em um DRIVE_FILE_ID real. Se o usuario pedir links e eles nao estiverem disponiveis, diga que nao encontrou.\n"
        "11. PESQUISA DE ACOES: se o usuario pedir para pesquisar/localizar uma acao ou tarefa, use consultar_historico_acoes primeiro. Nao substitua resultado ausente por acervo, email ou internet, salvo se o usuario pedir explicitamente essa ampliacao.\n"
        "12. ACESSO FINANCEIRO: para qualquer dado sobre rendas, contas, metas ou balanco interno, use consultar_financas_v2. Para novos registros, use obrigatoriamente propor_lancamento_financeiro para que o usuário receba os botões de confirmação. Detalhe os valores com precisao absoluta conforme retornado pelo sistema.\n"
        "13. EFICIENCIA: quando precisar de varias consultas independentes, solicite todas na mesma rodada de ferramentas. Evite rodadas sequenciais se uma unica rodada paralela resolver. Nunca chame mais de uma ferramenta de escrita/registro no mesmo turno; proponha uma confirmacao por vez.\n"
        "14. WHATSAPP: para enviar ou agendar mensagem de WhatsApp, use schedule_whatsapp_message. A ferramenta deve apenas preparar a proposta; o envio real depende de confirmacao por botao.\n"
    )

    if not acao_snapshot:
        return base

    titulo = acao_snapshot.get("titulo", "")
    plano_lines = "\n".join(f"  {p}" for p in acao_snapshot.get("plano_acao", [])) or "  (sem passos registrados)"
    diario_lines = "\n".join(f"  {d}" for d in acao_snapshot.get("acompanhamento_recente", [])) or "  (sem entradas recentes)"
    diario_integral = acao_snapshot.get("diario_integral", "") or "(sem diario integral)"
    arquivos = acao_snapshot.get("arquivos_disponiveis", []) or []
    arquivos_lines = (
        "\n".join(f"  - {a.get('nome', 'Arquivo sem nome')} | DRIVE_FILE_ID: {a.get('drive_file_id', '')}" for a in arquivos)
        or "  (sem arquivos disponiveis)"
    )
    tags_str = ", ".join(str(t) for t in acao_snapshot.get("tags", [])) or "nenhuma"

    acao_section = (
        "\n\n## ACAO VINCULADA - MODO CONTEXTO TRANCADO\n"
        f"ID: {acao_snapshot.get('id', '')}\n"
        f"Titulo: {titulo}\n"
        f"Area: {acao_snapshot.get('area_tematica', acao_snapshot.get('area', ''))}\n"
        f"Status: {acao_snapshot.get('status', '')}\n"
        f"Descricao: {acao_snapshot.get('descricao', '') or 'nao informada'}\n"
        f"Sintese: {acao_snapshot.get('sintese', '') or 'nao informada'}\n"
        f"Plano de Acao:\n{plano_lines}\n"
        f"Plano Atual Estruturado: {json.dumps(acao_snapshot.get('plano_atual', []), ensure_ascii=False)}\n"
        f"Diario Recente:\n{diario_lines}\n"
        f"Diario Integral:\n{diario_integral}\n"
        f"Notas: {acao_snapshot.get('notas', '') or 'nenhuma'}\n"
        f"Tags: {tags_str}\n\n"
        f"Arquivos Disponiveis:\n{arquivos_lines}\n\n"
        "MODO CONTEXTO TRANCADO ATIVADO - REGRAS ABSOLUTAS:\n"
        "1. O historico de conversas anteriores foi isolado. Concentre-se exclusivamente nesta acao.\n"
        "2. Use SOMENTE os dados desta secao ao descrever a acao. PROIBIDO usar RAG ou inferencia.\n"
        "3. Este payload replica a estrutura de contexto da interface web: plano atual, diario integral, tags e arquivos disponiveis.\n"
        "4. Ao receber perguntas sobre a acao, responda com base nos dados acima.\n"
        "5. Se o usuario pedir para atualizar, criar passos ou registrar progresso, use as ferramentas disponiveis.\n"
        "6. Se o usuario pedir links ou anexos, cite apenas as URLs/DRIVE_FILE_ID listados em Arquivos Disponiveis. Nunca use rotulos clicaveis como 'Acessar Pasta' sem mostrar a URL real.\n"
        "7. PROIBIDO usar consultar_historico_acoes neste modo — a acao ja esta carregada acima. Use apenas ferramentas de atualizacao (plano, diario, notas).\n"
    )
    return base + acao_section

def _normalize_for_matching(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text or "")
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return normalized.lower()

def _is_explicit_web_request(text: str) -> bool:
    lowered = _normalize_for_matching(text)
    web_markers = [
        "na internet",
        "na web",
        "pesquise na internet",
        "busque na internet",
        "pesquise no google",
        "busque no google",
        "noticia",
        "noticias",
        "site oficial",
        "link externo",
        "fonte externa",
        "fonte oficial",
    ]
    return any(marker in lowered for marker in web_markers)

def _extract_document_text(file_bytes: bytes, filename: str, mime_type: str) -> str | None:
    import io
    name = filename.lower()
    mime = mime_type.lower()
    
    # Word DOCX
    if mime == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' or name.endswith('.docx'):
        try:
            import mammoth
            result = mammoth.extract_raw_text(io.BytesIO(file_bytes))
            return (result.value or "").strip()
        except Exception as e:
            print(f"[Parser] Erro DOCX: {e}")
            
    # Excel XLSX / XLS
    elif mime in ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'application/vnd.ms-excel') or name.endswith(('.xlsx', '.xls')):
        try:
            import pandas as pd
            df_list = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
            sheets_text = []
            for sheet_name, df in df_list.items():
                sheets_text.append(f"ABA: {sheet_name}\n{df.to_csv(index=False)}")
            return "\n\n".join(sheets_text).strip()
        except Exception as e:
            print(f"[Parser] Erro Excel: {e}")
            
    # PowerPoint PPTX
    elif mime == 'application/vnd.openxmlformats-officedocument.presentationml.presentation' or name.endswith('.pptx'):
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            slides_text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, 'text') and shape.text.strip():
                        slides_text.append(shape.text.strip())
            return '\n'.join(slides_text).strip()
        except Exception as e:
            print(f"[Parser] Erro PPTX: {e}")
            
    # Plain text / CSV
    elif mime.startswith('text/') or name.endswith(('.txt', '.csv', '.json', '.xml', '.yaml', '.yml', '.md')):
        try:
            return file_bytes.decode('utf-8', errors='replace')
        except Exception as e:
            print(f"[Parser] Erro Texto: {e}")
            
    return None

def _is_internal_hermes_request(text: str) -> bool:
    lowered = _normalize_for_matching(text)
    internal_markers = [
        "hermes",
        "no sistema",
        "no copilot",
        "copiloto",
        "acao",
        "acoes",
        "tarefa",
        "tarefas",
        "agendamento",
        "documentacao",
        "deferido",
        "deferidos",
        "assistencia estudantil",
        "historico",
        "cadastro",
    ]
    return any(marker in lowered for marker in internal_markers)

def _extract_response_mode(text: str, session: dict) -> tuple[str, str]:
    raw_text = (text or "").strip()
    if not raw_text:
        return "texto", raw_text

    lowered = _normalize_for_matching(raw_text)
    explicit_audio_markers = (
        "/audio",
        "#audio",
        "[audio]",
        "mensagem de voz",
        "em voz",
        "por voz",
    )
    audio_keywords = ("audio", "voz")
    request_verbs = ("responda", "resposta", "mande", "manda", "envie", "quero")
    wants_audio = any(marker in lowered for marker in explicit_audio_markers)
    if not wants_audio:
        wants_audio = any(v in lowered for v in request_verbs) and any(k in lowered for k in audio_keywords)

    if wants_audio:
        cleaned = raw_text
        cleaned = re.sub(r"^/audio\b", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"#audio|\[audio\]", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"(me\s+)?responda\s+(em|por|para|pro?)\s+(o\s+)?[aá]udio", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"resposta\s+(em|por|para|pro?)\s+(o\s+)?[aá]udio", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"envie\s+(em|por|para|pro?)\s+(o\s+)?[aá]udio", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"(me\s+)?mande\s+(em|por|para|pro?)\s+(o\s+)?[aá]udio", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"quero\s+(em|por|para|pro?)\s+(o\s+)?[aá]udio", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"mensagem de voz", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"\b(em|por)\s+voz\b", "", cleaned, flags=re.IGNORECASE).strip()
        return "audio", cleaned.strip(" ,.-")

    explicit_text_markers = ("/texto", "#texto", "[texto]")
    if any(marker in lowered for marker in explicit_text_markers):
        cleaned = re.sub(r"^/texto\b|#texto|\[texto\]", "", raw_text, flags=re.IGNORECASE).strip()
        return "texto", cleaned

    return "texto", raw_text

def _extract_voice_profile(text: str, default_voice: str = "masculina") -> tuple[str, str]:
    raw_text = (text or "").strip()
    if not raw_text:
        return default_voice, raw_text

    lowered = _normalize_for_matching(raw_text)
    masculine_markers = (
        "#voz:masculina",
        "[voz:masculina]",
        "/voz masculina",
        "voz masculina",
        "na voz masculina",
        "com voz masculina",
    )

    if any(marker in lowered for marker in masculine_markers):
        cleaned = re.sub(r"#voz:masculina|\[voz:masculina\]|/voz masculina", "", raw_text, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"\b(na|com)?\s*voz masculina\b", "", cleaned, flags=re.IGNORECASE).strip()
        return "masculina", cleaned

    return default_voice, raw_text

def _run_gemini_text(gemini_key: str, system_instruction: str, user_prompt: str, model_id: str = _TEXT_MODEL_ID) -> str:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=gemini_key)
    response = generate_content_logged(
        client,
        model=model_id,
        contents=user_prompt,
        feature="telegram.tts_script",
        config=types.GenerateContentConfig(system_instruction=system_instruction),
    )
    return (response.text or "").strip()

def _build_tts_director_instruction(voice_profile: str) -> str:
    return (
        "Voce e o diretor de TTS do Hermes. "
        "Recebera uma resposta factual pronta e deve apenas converte-la em um roteiro curto para fala natural. "
        "Nao adicione fatos novos. "
        "Nao inclua links, nomes tecnicos longos nem listas densas. "
        "Nao use tags, colchetes nem anotacoes de palco. "
        "Entregue apenas o texto a ser falado. "
        "Mantenha tom de assistente operacional, conversa um a um, sem soar como narrador. "
        "Priorize objetividade, proximidade e sobriedade. "
        f"Perfil de voz desejado: {voice_profile}. "
        "Responda apenas com o roteiro final."
    )

def _run_gemini_tts(gemini_key: str, script_text: str, voice_profile: str) -> tuple[bytes, str]:
    from google import genai
    from google.genai import types

    voice_name = _DEFAULT_MALE_VOICE
    client = genai.Client(api_key=gemini_key)
    style_prompt = (
        "### DIRECTOR'S NOTES\n"
        "Style: concise operational assistant, one-to-one, practical and calm.\n"
        "Pacing: natural conversational pace, with short pauses only when meaning changes.\n"
        "Tone: professional, direct, grounded, helpful and discreet.\n"
        "Delivery: do not sound like a narrator, announcer, storyteller, presenter or commercial voice-over.\n"
        "Delivery: sound like a senior assistant speaking directly to one person.\n\n"
        "### SCRIPT\n"
        f"\"{script_text[:_MAX_TTS_TRANSCRIPT_CHARS]}\""
    )
    response = generate_content_logged(
        client,
        model=_TTS_MODEL_ID,
        contents=style_prompt,
        feature="telegram.tts_audio",
        config=types.GenerateContentConfig(
            response_modalities=["AUDIO"],
            speech_config=types.SpeechConfig(
                voice_config=types.VoiceConfig(
                    prebuilt_voice_config=types.PrebuiltVoiceConfig(voice_name=voice_name)
                )
            ),
        ),
    )

    for candidate in (response.candidates or []):
        for part in (candidate.content.parts or []):
            inline_data = getattr(part, "inline_data", None)
            if inline_data and getattr(inline_data, "data", None):
                return inline_data.data, getattr(inline_data, "mime_type", "audio/L16;rate=24000")
    raise RuntimeError("Gemini TTS nÃ£o retornou Ã¡udio.")

def _transcode_audio_for_telegram_voice(audio_bytes: bytes, mime_type: str) -> tuple[bytes, str, str] | tuple[None, None, None]:
    import shutil
    import subprocess

    sample_rate = 24000
    rate_match = re.search(r"rate=(\d+)", (mime_type or "").lower())
    if rate_match:
        sample_rate = int(rate_match.group(1))

    ffmpeg_bin = shutil.which("ffmpeg")
    if not ffmpeg_bin:
        return None, None, None

    with tempfile.TemporaryDirectory() as tmpdir:
        wav_path = os.path.join(tmpdir, "input.wav")
        ogg_path = os.path.join(tmpdir, "output.ogg")
        with wave.open(wav_path, "wb") as wav_file:
            wav_file.setnchannels(1)
            wav_file.setsampwidth(2)
            wav_file.setframerate(sample_rate)
            wav_file.writeframes(audio_bytes)
        cmd = [
            ffmpeg_bin,
            "-y",
            "-i",
            wav_path,
            "-c:a",
            "libopus",
            "-b:a",
            "32k",
            ogg_path,
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=120, check=False)
        if proc.returncode != 0 or not os.path.exists(ogg_path):
            print(f"[TTS] ffmpeg transcode failed: {proc.stderr[:400]}")
            return None, None, None
        with open(ogg_path, "rb") as f:
            return f.read(), "audio/ogg", "hermes_voice.ogg"

def _wrap_pcm_audio_as_wav(audio_bytes: bytes, mime_type: str) -> tuple[bytes, str, str]:
    sample_rate = 24000
    rate_match = re.search(r"rate=(\d+)", (mime_type or "").lower())
    if rate_match:
        sample_rate = int(rate_match.group(1))

    wav_buffer = BytesIO()
    with wave.open(wav_buffer, "wb") as wav_file:
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(sample_rate)
        wav_file.writeframes(audio_bytes)
    return wav_buffer.getvalue(), "audio/wav", "hermes_audio.wav"

def _send_tts_failure_notice(token: str, chat_id: str | int):
    _send_telegram_message(
        token,
        chat_id,
        "⚠️ Houve uma falha tÃ©cnica ao gerar o Ã¡udio. Repita a solicitaÃ§Ã£o ou peÃ§a a resposta em texto.",
    )

def _send_tts_failure_notice_contextual(db, token: str, chat_id: str | int, session: dict | None = None):
    _send_telegram_session_message(
        db,
        token,
        chat_id,
        "⚠️ Houve uma falha técnica ao gerar o áudio. Repita a solicitação ou peça a resposta em texto.",
        session=session,
    )

def _send_contextual_response(
    db,
    token: str,
    chat_id: str | int,
    text: str,
    session: dict | None = None,
    response_mode: str = "texto",
    gemini_key: str | None = None,
    voice_profile: str = "masculina",
    inline_keyboard: list | None = None,
    perf_state: dict | None = None,
):
    """Envia resposta factual pronta e, se pedido, converte o mesmo texto em audio."""
    message_id = _send_telegram_session_message(
        db,
        token,
        chat_id,
        text,
        session=session,
        inline_keyboard=inline_keyboard,
    )
    if response_mode != "audio":
        return message_id
    if not gemini_key:
        _send_tts_failure_notice_contextual(db, token, chat_id, session=session)
        return message_id

    try:
        with _telegram_action_heartbeat(token, chat_id, "record_voice"):
            tts_script = _run_gemini_text(
                gemini_key=gemini_key,
                system_instruction=_build_tts_director_instruction(voice_profile),
                user_prompt=text,
            )
            audio_bytes, audio_mime = _run_gemini_tts(
                gemini_key=gemini_key,
                script_text=tts_script or text,
                voice_profile=voice_profile,
            )
            tg_audio_bytes, tg_audio_mime, tg_audio_name = _transcode_audio_for_telegram_voice(audio_bytes, audio_mime)
            if tg_audio_bytes:
                if not _send_telegram_voice(token, chat_id, tg_audio_bytes, tg_audio_name, tg_audio_mime):
                    raise RuntimeError("Falha ao enviar voice note pelo Telegram.")
            else:
                wav_bytes, wav_mime, wav_name = _wrap_pcm_audio_as_wav(audio_bytes, audio_mime)
                if not _send_telegram_document(token, chat_id, wav_bytes, wav_name, wav_mime, caption="Resposta em audio"):
                    raise RuntimeError("Falha ao enviar arquivo de audio.")
        if perf_state is not None:
            _perf_mark(perf_state, "telegram.audio_response")
    except Exception as tts_err:
        print(f"[Core] Deterministic TTS error: {tts_err}")
        if perf_state is not None:
            _perf_log("telegram.request.deterministic_tts_error", perf_state, {"chat_id": str(chat_id), "error": str(tts_err)})
        _send_tts_failure_notice_contextual(db, token, chat_id, session=session)
    return message_id

def _run_gemini_turn(
    db,
    gemini_key: str,
    system_instruction: str,
    history: list,
    user_message_parts: list,
    tools_list: list,
    function_map: dict,
    perf_state: dict | None = None,
    sistema_id_contexto: str = None,
) -> str:
    """Runs a full Gemini multi-turn exchange and returns the final text response."""
    from concurrent.futures import ThreadPoolExecutor
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=gemini_key)
    model_id = _TEXT_MODEL_ID
    read_only_parallel_tools = {
        "consultar_historico_acoes",
        "buscar_arquivos_acervo",
        "pesquisar_internet",
        "ler_pagina_web",
        "consultar_agenda",
        "encontrar_slot_livre",
        "consultar_financas_v2",
        "buscar_e_analisar_email",
    }

    def _is_parallel_safe_tool(tool_name: str) -> bool:
        return tool_name in read_only_parallel_tools

    def _execute_tool_call(fc, *, parallel: bool = False):
        fn = function_map.get(fc.name)
        tool_start_ms = _perf_now_ms()
        if fn is None:
            result_text = f"Ferramenta '{fc.name}' não encontrada."
        else:
            try:
                kwargs = dict(fc.args or {})
                result = fn(**kwargs)
                result_text = result if isinstance(result, str) else json.dumps(result, ensure_ascii=False)
            except Exception as tool_err:
                result_text = f"Erro ao executar {fc.name}: {tool_err}"
        if perf_state is not None:
            perf_state.setdefault("tool_calls", []).append({
                "name": fc.name,
                "duration_ms": max(0, _perf_now_ms() - tool_start_ms),
                "parallel": parallel,
            })
        return types.Part.from_function_response(
            name=fc.name,
            response={"result": result_text},
        )

    def _tool_call_groups(func_calls: list) -> list[list]:
        groups = []
        current_read_group = []
        for fc in func_calls:
            if _is_parallel_safe_tool(fc.name):
                current_read_group.append(fc)
                continue
            if current_read_group:
                groups.append(current_read_group)
                current_read_group = []
            groups.append([fc])
        if current_read_group:
            groups.append(current_read_group)
        return groups

    clean_history = _sanitize_chat_history(history, types)
    # Schemas saneados: sem `default`, sem OBJECT vazio e com `required` derivado
    # da assinatura real (parâmetro com valor padrão não é obrigatório).
    declared_tools = _build_gemini_tools(client, types, tools_list)
    # Thinking de acordo com a família do modelo (3.x = thinking_level).
    thinking_config = _thinking_config_for_model(types, model_id)

    def _abrir_chat(*, com_thinking: bool, com_historico: bool):
        config_kwargs = {
            "system_instruction": system_instruction,
            "tools": declared_tools,
            "automatic_function_calling": types.AutomaticFunctionCallingConfig(disable=True),
        }
        # Sem thinking: evita parts com thought=True que confundem a extração de texto.
        if com_thinking and thinking_config is not None:
            config_kwargs["thinking_config"] = thinking_config
        return client.chats.create(
            model=model_id,
            config=types.GenerateContentConfig(**config_kwargs),
            history=clean_history if com_historico else [],
        )

    # Degradação diagnosticável: a API devolve 400 INVALID_ARGUMENT genérico, sem
    # apontar o campo. Cada tentativa remove um suspeito e registra no log qual
    # delas passou — assim uma recaída fica identificada na primeira ocorrência,
    # em vez de virar só "⚠️ Erro ao processar" para o usuário.
    tentativas = (
        ("completa", True, True),
        ("sem thinking_config", False, True),
        ("sem histórico", False, False),
    )
    response = None
    erro_invalid_argument = None
    for rotulo, com_thinking, com_historico in tentativas:
        try:
            chat = _abrir_chat(com_thinking=com_thinking, com_historico=com_historico)
            if perf_state is not None:
                _perf_mark(perf_state, "telegram.chat_create")
            response = send_message_logged(
                chat,
                user_message_parts,
                model=model_id,
                feature="telegram.first_turn",
                db=db,
            )
            if rotulo != "completa":
                print(f"[Core] Gemini aceitou o pedido na tentativa '{rotulo}' (modelo {model_id}).")
            break
        except Exception as send_err:
            if not _is_invalid_argument_error(send_err):
                raise
            erro_invalid_argument = send_err
            print(f"[Core] 400 INVALID_ARGUMENT na tentativa '{rotulo}' (modelo {model_id}): {send_err}")

    if response is None:
        raise erro_invalid_argument

    if perf_state is not None:
        _perf_mark(perf_state, "telegram.first_model_response")

    # Agentic loop — processa tool calls até obter resposta de texto
    for _ in range(6):
        if not response.candidates:
            break
        candidate = response.candidates[0]
        if candidate.finish_reason and candidate.finish_reason.name not in ("STOP", "MAX_TOKENS", ""):
            break

        # Coleta function calls neste turno
        func_calls = []
        for part in (candidate.content.parts or []):
            if hasattr(part, "function_call") and part.function_call and part.function_call.name:
                func_calls.append(part.function_call)

        if not func_calls:
            break

        # Executa todas as tool calls e coleta resultados
        tool_results = []
        for group in _tool_call_groups(func_calls):
            if len(group) == 1:
                tool_results.append(_execute_tool_call(group[0], parallel=False))
                continue

            max_workers = min(len(group), 6)
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = [executor.submit(_execute_tool_call, fc, parallel=True) for fc in group]
                for future in futures:
                    tool_results.append(future.result())

        if perf_state is not None:
            parallel_count = sum(1 for fc in func_calls if _is_parallel_safe_tool(fc.name))
            serial_count = len(func_calls) - parallel_count
            perf_state.setdefault("tool_rounds", []).append({
                "total_calls": len(func_calls),
                "parallel_safe_calls": parallel_count,
                "serialized_calls": serial_count,
            })
        response = send_message_logged(
            chat,
            tool_results,
            model=model_id,
            feature="telegram.tool_roundtrip",
            db=db,
        )
        if perf_state is not None:
            _perf_mark(perf_state, "telegram.tool_roundtrip")
    else:
        # Se saiu do loop por limite de iterações, força um turno final de texto
        last_chance_msg = [types.Part(text="Limite de pesquisas atingido. Por favor, apresente uma resposta final ao usuário com base no que você encontrou (ou informe que não encontrou).")]
        response = send_message_logged(
            chat,
            last_chance_msg,
            model=model_id,
            feature="telegram.safety_final_turn",
            db=db,
        )
        if perf_state is not None:
            _perf_mark(perf_state, "telegram.safety_final_turn")

    # Extrai texto final — ignora parts de thinking (thought=True) que não são resposta
    text_parts = []
    if response.candidates:
        for part in (response.candidates[0].content.parts or []):
            # Filtra parts de raciocínio interno (thinking) — não são texto de resposta
            if getattr(part, "thought", False):
                continue
            if hasattr(part, "text") and part.text:
                text_parts.append(part.text)
    return "\n".join(text_parts).strip() or "Peço desculpas, mas não consegui formular uma resposta. Tente reformular sua pergunta."

def _health_red_flag_active(data: dict) -> bool:
    radicular = data.get("radicular") or {}
    pain = data.get("pain") or {}
    if radicular.get("location") == "pe":
        return True
    if radicular.get("motorWeakness") is True:
        return True
    if (pain.get("morning") or 0) >= 9 or (pain.get("evening") or 0) >= 9:
        return True
    return False
